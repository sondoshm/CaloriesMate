# -*- coding: utf-8 -*-
"""
CalorieMate Pro v2
ميزات النسخة الجديدة:
 1) إصلاح حفظ الأكلات المخصصة
 2) تتبع الكوليسترول فعلياً مع شريط يتحول للأحمر
 3) سجل الأيام + رسم بياني (Canvas بدون مكتبات)
 4) تصفير تلقائي كل يوم مع الأرشفة
 5) تنبيهات الحديد والماء والكوليسترول
 6) وجبات مفضلة
 7) تعديل/حذف الأكلات المخصصة
 8) تتبع الوزن و BMI
 9) بحث ذكي (يتجاهل الهمزات والتشكيل) + RTL + وضع فاتح/داكن + خطة أسبوعية متغيرة
10) تصدير TXT / CSV / Excel / HTML (للطباعة كـ PDF)
"""
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
import json
import os
import sys
import csv
import re
import html
import webbrowser
from datetime import datetime, timedelta

# ==========================================
# قاعدة البيانات (لكل 100 جرام)
# (السعرات, بروتين, كارب, دهون, حديد_ملجم, وزن_الحصة_التقريبي)
# ==========================================
FOOD_DATABASE = {
    # --- الفواكه ---
    "جوافة": (68, 2.6, 14.3, 1.0, 0.3, 100),
    "تفاح": (52, 0.3, 13.8, 0.2, 0.1, 150),
    "موز": (89, 1.1, 22.8, 0.3, 0.3, 120),
    "برتقال": (47, 0.9, 11.8, 0.1, 0.1, 130),
    "أفوكادو": (160, 2.0, 8.5, 14.7, 0.6, 100),
    "بلح طازج (رطب)": (142, 1.8, 37.0, 0.2, 1.0, 50),
    "تمر (مجهول/خلاص)": (282, 2.5, 75.0, 0.4, 1.0, 30),
    "مشمش طازج": (48, 1.4, 11.0, 0.4, 0.4, 100),
    "مشمش مجفف (قطين)": (241, 3.4, 63.0, 0.5, 2.7, 40),
    "قراصيا (برقوق مجفف)": (240, 2.2, 64.0, 0.4, 2.5, 40),
    "بطيخ": (30, 0.6, 7.6, 0.2, 0.2, 200),
    "شمام": (34, 0.8, 8.2, 0.2, 0.2, 150),
    "أناناس": (50, 0.5, 13.1, 0.1, 0.3, 100),
    "كمثرى (إجاص)": (57, 0.4, 15.0, 0.1, 0.2, 150),
    "بابايا": (43, 0.5, 11.0, 0.3, 0.3, 140),
    "كرز طازج": (63, 1.1, 16.0, 0.2, 0.4, 100),
    "برقوق (خوخ أحمر)": (46, 0.7, 11.4, 0.3, 0.2, 100),
    "تين طازج": (74, 0.8, 19.2, 0.3, 0.4, 50),
    "تين مجفف": (249, 3.3, 63.9, 0.9, 2.0, 40),
    "عنب": (69, 0.7, 18.1, 0.2, 0.4, 100),
    "فراولة": (32, 0.7, 7.7, 0.3, 0.4, 100),
    "توت عليق (أسود)": (43, 1.4, 10.0, 0.5, 0.6, 100),
    "توت أزرق (Blueberry)": (57, 0.7, 14.0, 0.3, 0.3, 100),
    "مانجو": (60, 0.8, 15.0, 0.4, 0.2, 150),
    "كيوي": (61, 1.1, 14.7, 0.5, 0.3, 75),
    "خوخ دُراق": (39, 0.9, 9.5, 0.3, 0.3, 130),
    "رمان": (83, 1.7, 18.7, 1.2, 0.3, 150),
    "زبيب": (299, 3.1, 79.2, 0.5, 1.9, 40),
    "جريب فروت": (42, 0.8, 10.6, 0.1, 0.1, 120),
    "جوز هند مبشور": (660, 6.9, 23.7, 64.5, 3.3, 15),

    # --- المكسرات والبذور ---
    "بذور السمسم": (573, 17.7, 23.4, 49.7, 14.6, 15),
    "طحينية (عجينة السمسم)": (595, 17.0, 21.2, 53.8, 9.0, 15),
    "بذور القرع (بذور اليقطين)": (559, 30.0, 10.7, 49.0, 8.8, 30),
    "لوز خام": (579, 21.0, 22.0, 50.0, 3.7, 30),
    "لوز محمص": (598, 21.0, 21.0, 52.5, 3.8, 30),
    "جوز (عين جمل)": (654, 15.0, 14.0, 65.0, 2.9, 30),
    "فستق سوداني": (567, 25.8, 16.1, 49.2, 4.6, 30),
    "بذور الشمس": (584, 20.8, 20.0, 51.5, 5.3, 30),
    "كاجو": (553, 18.2, 30.2, 43.8, 6.7, 30),
    "فستق حلبي": (560, 20.0, 27.0, 45.0, 3.9, 30),
    "بذور الشيا": (486, 16.5, 42.1, 30.7, 7.7, 15),
    "بذور الكتان": (534, 18.3, 28.9, 42.2, 5.7, 15),
    "فشار بدون زيت": (387, 12.9, 77.9, 4.5, 3.2, 40),
    "فشار بالزيت/الزبدة": (535, 9.0, 57.0, 31.0, 2.8, 50),

    # --- الخضراوات ---
    "سبانخ": (23, 2.9, 3.6, 0.4, 2.7, 100),
    "خيار": (15, 0.7, 3.6, 0.1, 0.3, 100),
    "طماطم (بندورة)": (18, 0.9, 3.9, 0.2, 0.3, 120),
    "خس": (15, 1.4, 2.9, 0.2, 0.9, 50),
    "باذنجان": (25, 1.0, 5.9, 0.2, 0.2, 150),
    "بروكلي": (34, 2.8, 6.6, 0.4, 0.7, 100),
    "زهرة (قرنابيط)": (25, 1.9, 5.0, 0.3, 0.4, 100),
    "بطاطا مسلوقة": (87, 1.9, 20.1, 0.1, 0.8, 150),
    "بطاطا مقلية": (312, 3.4, 41.0, 15.0, 0.8, 150),
    "بطاطا حلوة مشوية": (90, 2.0, 20.7, 0.1, 0.6, 150),
    "جزر": (41, 0.9, 9.6, 0.2, 0.3, 80),
    "فلفل رومي أخضر/أحمر": (20, 0.9, 4.6, 0.2, 0.4, 100),
    "كوسا": (17, 1.2, 3.1, 0.3, 0.4, 100),

    # --- النشويات ---
    "أرز أبيض مطبوخ": (130, 2.7, 28.0, 0.3, 0.2, 150),
    "أرز بني مطبوخ": (111, 2.6, 23.0, 0.9, 0.5, 150),
    "أرز بسمتي مطبوخ": (121, 3.5, 25.0, 0.4, 0.3, 150),
    "برغل مطبوخ": (83, 3.1, 18.6, 0.2, 1.0, 150),
    "فريكة مطبوخة": (120, 4.5, 22.0, 0.8, 1.5, 150),
    "معكرونة مسلوقة": (131, 5.0, 25.0, 1.1, 1.3, 150),
    "خبز أبيض عربي": (265, 9.0, 53.0, 1.2, 1.5, 80),
    "خبز أسمر (قمح كامل)": (247, 13.0, 41.0, 3.4, 2.5, 80),
    "توست أبيض": (265, 8.8, 49.0, 3.2, 1.4, 30),
    "توست أسمر": (247, 12.0, 43.0, 3.5, 2.4, 30),
    "شوفان (خام)": (389, 16.9, 66.3, 6.9, 4.7, 40),

    # --- اللحوم والدواجن والأسماك ---
    "كبدة دجاج / كبدة بقر": (167, 24.5, 0.0, 6.5, 11.6, 120),
    "لحم بقر صافي (مشوي)": (250, 26.0, 0.0, 15.0, 2.6, 150),
    "لحم مفروم مطبوخ": (260, 24.0, 0.0, 18.0, 2.4, 120),
    "لحم غنم مشوي": (294, 25.0, 0.0, 21.0, 1.9, 150),
    "صدر دجاج مشوي": (165, 31.0, 0.0, 3.6, 1.0, 150),
    "فخذ دجاج مشوي": (209, 26.0, 0.0, 10.9, 1.3, 130),
    "سمك مشوي": (105, 22.0, 0.0, 2.5, 0.8, 180),
    "سمك سلمون مشوي": (206, 22.0, 0.0, 12.0, 0.8, 150),
    "سمك تونة بالماء": (116, 26.0, 0.0, 1.0, 1.3, 120),
    "سمك تونة بالزيت": (198, 29.0, 0.0, 8.2, 1.4, 120),
    "جمبري (روبيان)": (99, 24.0, 0.2, 0.3, 2.4, 120),

    # --- الألبان والبيض ---
    "بيض (مسلوق)": (155, 12.6, 1.1, 10.6, 1.8, 50),
    "بيض (مقلي بالزيت)": (196, 13.6, 0.8, 15.3, 1.9, 60),
    "جبنة بيضاء بلدية": (210, 15.0, 1.5, 16.0, 0.3, 50),
    "جبنة شيدر": (403, 25.0, 1.3, 33.0, 0.7, 30),
    "جبنة موزاريلا": (280, 28.0, 3.1, 17.0, 0.4, 40),
    "جبنة قريش": (98, 11.0, 3.4, 4.3, 0.1, 100),
    "حليب كامل الدسم": (61, 3.2, 4.8, 3.3, 0.1, 240),
    "حليب قليل الدسم": (42, 3.4, 5.0, 1.0, 0.1, 240),
    "لبن زبادي كامل الدسم": (61, 3.5, 4.7, 3.3, 0.1, 170),
    "لبن زبادي يوناني": (59, 10.0, 3.6, 0.4, 0.1, 170),
    "لبنة": (150, 6.0, 4.0, 12.0, 0.2, 50),

    # --- البقوليات ---
    "عدس مطبوخ": (116, 9.0, 20.0, 0.4, 3.3, 160),
    "حمص مسلوق": (164, 8.9, 27.0, 2.6, 2.9, 150),
    "فول مدمس": (110, 7.6, 19.8, 0.4, 2.5, 160),
    "فاصوليا بيضاء مطبوخة": (140, 9.7, 25.0, 0.5, 3.7, 160),
    "ترمس مسلوق": (119, 15.6, 9.9, 2.8, 1.2, 100),

    # --- الأكلات الشعبية ---
    "حمص بطحينة": (166, 8.0, 14.0, 9.6, 2.5, 100),
    "شاورما دجاج (لحم فقط)": (220, 25.0, 2.0, 12.0, 1.5, 150),
    "شاورما لحم (لحم فقط)": (250, 23.0, 2.0, 16.0, 2.8, 150),
    "فلافل": (333, 13.3, 31.8, 17.8, 3.4, 60),
    "كبسة بالدجاج": (160, 11.0, 20.0, 4.5, 1.1, 250),
    "كبسة باللحم": (190, 12.0, 19.0, 7.5, 1.8, 250),
    "مجدرة": (140, 6.0, 22.0, 3.5, 2.1, 200),
    "مسخن دجاج": (210, 14.0, 15.0, 11.0, 1.6, 250),
    "منسف": (195, 12.0, 18.0, 8.5, 1.7, 300),
    "ملوخية مع دجاج": (95, 8.0, 5.0, 4.5, 2.3, 200),
}

# الكوليسترول (ملغم لكل 100 جرام) - قيم تقريبية. أي صنف غير مذكور = 0 (نباتي)
CHOL_100G = {
    "كبدة دجاج / كبدة بقر": 350,
    "لحم بقر صافي (مشوي)": 90,
    "لحم مفروم مطبوخ": 85,
    "لحم غنم مشوي": 95,
    "صدر دجاج مشوي": 85,
    "فخذ دجاج مشوي": 93,
    "سمك مشوي": 60,
    "سمك سلمون مشوي": 63,
    "سمك تونة بالماء": 30,
    "سمك تونة بالزيت": 35,
    "جمبري (روبيان)": 189,
    "بيض (مسلوق)": 373,
    "بيض (مقلي بالزيت)": 400,
    "جبنة بيضاء بلدية": 60,
    "جبنة شيدر": 105,
    "جبنة موزاريلا": 79,
    "جبنة قريش": 17,
    "حليب كامل الدسم": 10,
    "حليب قليل الدسم": 5,
    "لبن زبادي كامل الدسم": 13,
    "لبن زبادي يوناني": 5,
    "لبنة": 40,
    "شاورما دجاج (لحم فقط)": 80,
    "شاورما لحم (لحم فقط)": 85,
    "كبسة بالدجاج": 35,
    "كبسة باللحم": 40,
    "مسخن دجاج": 45,
    "منسف": 40,
    "ملوخية مع دجاج": 30,
}

# ==========================================
# الثوابت والثيمات
# ==========================================
FONT = "Segoe UI"
FROZEN = getattr(sys, "frozen", False)
# مكان ملفات البرنامج نفسها (الأيقونة وغيرها)
APP_DIR = getattr(sys, "_MEIPASS", None) or os.path.dirname(os.path.abspath(__file__))
if FROZEN:
    # بعد التثبيت: البيانات بمجلد AppData (يكتب فيه ويندوز بدون صلاحيات) والتقارير بـ Documents
    BASE_DIR = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "CalorieMate")
    REPORT_DIR = os.path.join(os.path.expanduser("~"), "Documents", "CalorieMate Reports")
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    REPORT_DIR = BASE_DIR
for _d in (BASE_DIR, REPORT_DIR):
    try:
        os.makedirs(_d, exist_ok=True)
    except OSError:
        pass
DATA_FILE = os.path.join(BASE_DIR, "calories_data.json")
BACKUP_DIR = os.path.join(BASE_DIR, "backups")
if FROZEN and not os.path.exists(DATA_FILE):
    # ترحيل بياناتك القديمة لو كانت جنب ملف الـ exe
    _old = os.path.join(os.path.dirname(sys.executable), "calories_data.json")
    if os.path.exists(_old):
        try:
            import shutil
            shutil.copy2(_old, DATA_FILE)
        except OSError:
            pass


def set_icon(win):
    try:
        win.iconbitmap(os.path.join(APP_DIR, "icon.ico"))
    except Exception:
        pass

THEMES = {
    # ألوان ناعمة ومريحة للعين: لا أسود صافي ولا أبيض صافي
    "dark": {
        "MAIN_BG": "#141B26", "CARD_BG": "#1F2937", "ACCENT": "#2F8F83",
        "TEXT": "#E6EBF2", "MUTED": "#8D9BB0", "GOLD": "#E8B66E",
        "DANGER": "#E5766A", "IRON": "#EC7A8A", "WATER": "#6DB3E3",
        "GOOD": "#3E9E8A", "BLUE": "#85B8E8",
    },
    "light": {
        "MAIN_BG": "#EFEBE4", "CARD_BG": "#FBF9F5", "ACCENT": "#2F8F83",
        "TEXT": "#2B3440", "MUTED": "#6B7684", "GOLD": "#B06A1E",
        "DANGER": "#CF5546", "IRON": "#C2455A", "WATER": "#2D80B5",
        "GOOD": "#2C8873", "BLUE": "#2D80B5",
    },
}
C = dict(THEMES["dark"])

MEAL_TYPES = ("فطور", "غداء", "عشاء", "سناك")
TYPE_ICONS = {"فطور": "🌅", "غداء": "☀️", "عشاء": "🌙", "سناك": "🍎"}
# مفهرسة حسب weekday(): الاثنين=0 ... الأحد=6
DAY_NAMES = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
DAY_ORDER = [5, 6, 0, 1, 2, 3, 4]  # ترتيب العرض: السبت أولاً

WEEK_PLAN = {
    0: ("شوفان بالحليب قليل الدسم + ملعقة بذور كتان + برتقالة",
        "شوربة عدس بالليمون + سلطة سبانخ وفلفل رومي + خبز أسمر",
        "صدر دجاج مشوي + بروكلي وزهرة مشوية + أرز بني",
        "حفنة لوز نيئ + تفاحة"),
    1: ("فول مدمس بالليمون وقليل من زيت الزيتون + خيار وطماطم + خبز أسمر",
        "مجدرة + سلطة خضراء + لبن زبادي قليل الدسم",
        "سمك مشوي + خضار سوتيه + برغل",
        "3 حبات تمر + حفنة جوز"),
    2: ("توست أسمر + جبنة قريش + طماطم وخس + جوافة",
        "حمص مسلوق مع برغل وسلطة سبانخ بالليمون",
        "تونة بالماء مع سلطة خضراء وخبز أسمر",
        "لبن يوناني + توت أو فراولة"),
    3: ("شوفان بالحليب + مشمش مجفف + بذور شيا",
        "ملوخية مع صدر دجاج + أرز بني",
        "شوربة خضار + فاصوليا بيضاء مطبوخة + سلطة",
        "بذور قرع + برتقالة"),
    4: ("بيضة مسلوقة واحدة + خبز أسمر + خيار وطماطم",
        "لحم قليل الدهن مشوي (100 جرام) + سلطة سبانخ بالليمون + برغل",
        "عدس مطبوخ مع خضار + لبن زبادي",
        "تين مجفف حبتين + لوز"),
    5: ("شوفان بالموز وبذور الكتان + حليب قليل الدسم",
        "كبسة بالدجاج (صدر بدون جلد) + سلطة كبيرة",
        "سلطة حمص وخضار + خبز أسمر",
        "فشار بدون زيت + تفاحة"),
    6: ("فطيرة سبانخ بالفرن بخبز أسمر + لبن زبادي",
        "منسف خفيف (لحم قليل + أرز كثير + سلطة) أو فاصوليا بيضاء",
        "صدر دجاج مشوي + كوسا وجزر مشوي",
        "قراصيا حبتين + حفنة لوز"),
}

# وحدات بيتية تقريبية (جرام) لبعض الأصناف
MEASURES = {
    "بيض (مسلوق)": {"حبة": 50}, "بيض (مقلي بالزيت)": {"حبة": 55},
    "تفاح": {"حبة": 180}, "موز": {"حبة": 120}, "برتقال": {"حبة": 130}, "جوافة": {"حبة": 100},
    "كيوي": {"حبة": 75}, "كمثرى (إجاص)": {"حبة": 160}, "خوخ دُراق": {"حبة": 130},
    "مشمش طازج": {"حبة": 35}, "مشمش مجفف (قطين)": {"حبة": 8}, "قراصيا (برقوق مجفف)": {"حبة": 10},
    "تمر (مجهول/خلاص)": {"حبة": 12}, "بلح طازج (رطب)": {"حبة": 10},
    "تين طازج": {"حبة": 50}, "تين مجفف": {"حبة": 15},
    "خيار": {"حبة": 120}, "طماطم (بندورة)": {"حبة": 120}, "جزر": {"حبة": 80},
    "كوسا": {"حبة": 100}, "فلفل رومي أخضر/أحمر": {"حبة": 120}, "بطاطا مسلوقة": {"حبة": 150},
    "حليب كامل الدسم": {"كاسة": 240}, "حليب قليل الدسم": {"كاسة": 240},
    "لبن زبادي كامل الدسم": {"علبة": 170, "كاسة": 240}, "لبن زبادي يوناني": {"علبة": 170},
    "خبز أبيض عربي": {"رغيف": 80}, "خبز أسمر (قمح كامل)": {"رغيف": 80},
    "توست أبيض": {"شريحة": 30}, "توست أسمر": {"شريحة": 30},
    "أرز أبيض مطبوخ": {"كاسة": 160}, "أرز بني مطبوخ": {"كاسة": 160}, "أرز بسمتي مطبوخ": {"كاسة": 160},
    "برغل مطبوخ": {"كاسة": 150}, "عدس مطبوخ": {"كاسة": 190}, "حمص مسلوق": {"كاسة": 165},
    "معكرونة مسلوقة": {"كاسة": 140}, "شوفان (خام)": {"كاسة": 80},
    "لوز خام": {"حفنة": 28}, "لوز محمص": {"حفنة": 28}, "جوز (عين جمل)": {"حفنة": 28},
    "كاجو": {"حفنة": 28}, "فستق حلبي": {"حفنة": 28}, "فستق سوداني": {"حفنة": 28},
    "جبنة بيضاء بلدية": {"قطعة": 30}, "صدر دجاج مشوي": {"قطعة": 120}, "فخذ دجاج مشوي": {"قطعة": 130},
    "سمك مشوي": {"قطعة": 150}, "فلافل": {"حبة": 17},
}

# أطعمة غنية بفيتامين C (لتلميح امتصاص الحديد)
VITC_FOODS = {
    "جوافة", "برتقال", "فراولة", "كيوي", "جريب فروت", "فلفل رومي أخضر/أحمر", "بروكلي",
    "طماطم (بندورة)", "أناناس", "مانجو", "بابايا", "زهرة (قرنابيط)", "توت أزرق (Blueberry)", "رمان",
}

_AR_TRANS = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ة": "ه", "ى": "ي", "ـ": None})
_TASHKEEL = re.compile("[\u064B-\u0652\u0670]")


def norm(text):
    """تطبيع النص العربي للبحث: يتجاهل التشكيل والهمزات والتاء المربوطة."""
    return _TASHKEEL.sub("", text.strip().lower()).translate(_AR_TRANS)


def today_str():
    return datetime.now().strftime("%Y-%m-%d")


def fmt(v):
    return str(int(round(v))) if abs(v) >= 10 else str(round(v, 1))


def totals(meals):
    return {k: sum(m.get(k, 0) for m in meals) for k in ("cal", "p", "c", "f", "fe", "chol")}


def new_today():
    return {"date": today_str(), "meals": [], "water": 0.0,
            "iron_notified": False, "chol_warned": False, "supp_taken": []}


def default_data():
    return {
        "daily_goal": 2000, "water_goal": 2.5, "theme": "dark",
        "user_profile": {"gender": "أنثى", "weight": "65", "height": "165", "age": "30",
                         "activity": "قليل النشاط (مكتبي)",
                         "cholesterol": "مرتفع (حد أقصى 200 mg)"},
        "settings": {"water_reminder_hrs": 2.0, "p_pct": 25, "c_pct": 50, "f_pct": 25, "vitc_tip": True},
        "custom_foods": {}, "favorites": {}, "weights": [], "history": {},
        "supplements": [], "usage": {}, "last_summary": "",
        "today": new_today(),
    }


def parse_data(raw):
    d = default_data()
    for k in ("daily_goal", "water_goal", "theme", "favorites", "weights", "history",
              "custom_foods", "supplements", "usage", "last_summary"):
        if k in raw:
            d[k] = raw[k]
    d["user_profile"].update(raw.get("user_profile", {}))
    d["settings"].update(raw.get("settings", {}))
    if "today" in raw:
        d["today"].update(raw["today"])
    else:  # ترحيل من النسخة القديمة
        d["today"]["meals"] = raw.get("meals", [])
        d["today"]["water"] = raw.get("water_intake", 0.0)
    return d


def load_data():
    if not os.path.exists(DATA_FILE):
        return default_data()
    try:
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            raw = json.load(f)
        if not isinstance(raw, dict):
            raise ValueError
        return parse_data(raw)
    except Exception:
        try:
            os.replace(DATA_FILE, DATA_FILE + ".bak")
        except OSError:
            pass
        return default_data()


def roll_day(d):
    """إذا تغير التاريخ: أرشفة اليوم القديم وبدء يوم جديد. ترجع True إذا صار تغيير."""
    cur = d["today"]
    if cur.get("date") == today_str():
        return False
    if cur.get("meals") or cur.get("water", 0) > 0:
        d["history"][cur["date"]] = {"meals": cur["meals"], "water": cur.get("water", 0.0)}
    d["today"] = new_today()
    return True


# ==========================================
# عناصر واجهة مساعدة
# ==========================================
def lbl(parent, text="", size=9, bold=False, fg=None, bg=None, **kw):
    return tk.Label(parent, text=text, font=(FONT, size, "bold" if bold else "normal"),
                    fg=fg or C["TEXT"], bg=bg or C["CARD_BG"], **kw)


def shade(hex_color, f):
    h = hex_color.lstrip("#")
    if len(h) != 6:
        return hex_color
    try:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return hex_color
    r, g, b = (max(0, min(255, int(x * f))) for x in (r, g, b))
    return f"#{r:02x}{g:02x}{b:02x}"


def lum(hex_color):
    h = hex_color.lstrip("#")
    if len(h) != 6:
        return 0
    try:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return 0
    return 0.299 * r + 0.587 * g + 0.114 * b


def btn(parent, text, cmd, bg, fg="white", size=9, bold=True, **kw):
    hover = shade(bg, 1.2 if lum(bg) < 128 else 0.93)
    b = tk.Button(parent, text=text, command=cmd, font=(FONT, size, "bold" if bold else "normal"),
                  bg=bg, fg=fg, activebackground=hover, activeforeground=fg,
                  bd=0, cursor="hand2", **kw)
    b.bind("<Enter>", lambda e: b.config(bg=hover) if str(b["state"]) != "disabled" else None)
    b.bind("<Leave>", lambda e: b.config(bg=bg))
    return b


def ent(parent, justify="right", bg=None, **kw):
    return tk.Entry(parent, font=(FONT, 10), bg=bg or C["MAIN_BG"], fg=C["TEXT"],
                    insertbackground=C["TEXT"], bd=0, justify=justify,
                    highlightthickness=1, highlightbackground=C["CARD_BG"],
                    highlightcolor=C["ACCENT"], **kw)


def optmenu(parent, var, *values, command=None):
    m = tk.OptionMenu(parent, var, *values, command=command)
    m.config(bg=C["MAIN_BG"], fg=C["TEXT"], activebackground=C["MAIN_BG"],
             activeforeground=C["TEXT"], bd=0, highlightthickness=0, font=(FONT, 9))
    m["menu"].config(bg=C["MAIN_BG"], fg=C["TEXT"])
    return m


def make_scrollable(parent):
    canvas = tk.Canvas(parent, bg=C["CARD_BG"], highlightthickness=0)
    sb = tk.Scrollbar(parent, orient="vertical", command=canvas.yview)
    inner = tk.Frame(canvas, bg=C["CARD_BG"])
    win_id = canvas.create_window((0, 0), window=inner, anchor="nw")
    inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.bind("<Configure>", lambda e: canvas.itemconfig(win_id, width=e.width))
    canvas.configure(yscrollcommand=sb.set)
    sb.pack(side="right", fill="y")
    canvas.pack(side="left", fill="both", expand=True)
    canvas.bind("<Enter>", lambda e: canvas.bind_all(
        "<MouseWheel>", lambda ev: canvas.yview_scroll(int(-ev.delta / 120), "units")))
    canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))
    return inner


def draw_bars(canvas, labels, values, limit, color, w, h):
    canvas.delete("all")
    pl, pb, pt = 12, 28, 22
    maxv = max(max(values), limit or 0, 1) * 1.15
    n = len(values)
    bw = (w - pl - 10) / n
    base = h - pb
    for i, v in enumerate(values):
        x0 = pl + i * bw + bw * 0.18
        x1 = pl + (i + 1) * bw - bw * 0.18
        y0 = base - (v / maxv) * (h - pb - pt)
        col = C["DANGER"] if (limit and v > limit) else color
        canvas.create_rectangle(x0, y0, x1, base, fill=col, outline="")
        if v > 0:
            canvas.create_text((x0 + x1) / 2, y0 - 8, text=fmt(v), fill=C["TEXT"], font=(FONT, 8, "bold"))
        canvas.create_text((x0 + x1) / 2, base + 12, text=labels[i], fill=C["MUTED"], font=(FONT, 8))
    canvas.create_line(pl, base, w - 5, base, fill=C["MUTED"])
    if limit:
        y = base - (limit / maxv) * (h - pb - pt)
        canvas.create_line(pl, y, w - 5, y, fill=C["GOLD"], dash=(4, 3))


def draw_line(canvas, labels, values, w, h):
    canvas.delete("all")
    pl, pb, pt = 30, 26, 20
    if not values:
        canvas.create_text(w / 2, h / 2, text="لا توجد قياسات بعد", fill=C["MUTED"], font=(FONT, 10))
        return
    lo, hi = min(values), max(values)
    if hi - lo < 1:
        lo, hi = lo - 1, hi + 1
    n = len(values)
    pts = []
    for i, v in enumerate(values):
        x = w / 2 if n == 1 else pl + i * (w - pl - 15) / (n - 1)
        y = (h - pb) - ((v - lo) / (hi - lo)) * (h - pb - pt)
        pts.append((x, y))
    if n > 1:
        canvas.create_line(*[c for p in pts for c in p], fill=C["WATER"], width=2)
    for (x, y), v, lab in zip(pts, values, labels):
        canvas.create_oval(x - 4, y - 4, x + 4, y + 4, fill=C["GOLD"], outline="")
        canvas.create_text(x, y - 12, text=fmt(v), fill=C["TEXT"], font=(FONT, 8, "bold"))
        canvas.create_text(x, h - pb + 12, text=lab, fill=C["MUTED"], font=(FONT, 7))


def open_path(path):
    try:
        if hasattr(os, "startfile"):
            os.startfile(path)
        else:
            webbrowser.open("file://" + path)
    except Exception:
        pass


# ==========================================
# التطبيق
# ==========================================
class App:
    def __init__(self):
        self.d = load_data()
        C.update(THEMES.get(self.d.get("theme", "dark"), THEMES["dark"]))
        roll_day(self.d)
        self.root = tk.Tk()
        self.root.title("CalorieMate Pro 🍎")
        set_icon(self.root)
        self.root.geometry("570x760")
        self.root.resizable(True, True)
        self.root.minsize(570, 760)
        self.last_water = datetime.now()
        self.banner_job = None
        self.supp_reminded = set()
        self.save()
        self.make_backup(auto=True)
        self.build_main()
        if datetime.now().weekday() == 6 and self.d.get("last_summary") != today_str():
            self.d["last_summary"] = today_str()
            self.save()
            self.show_banner("📈 ملخصك الأسبوعي جاهز! افتح «الملخص»", C["BLUE"], 20000)
        self.root.after(3000, self.tick)

    def run(self):
        self.root.mainloop()

    # ---------- بيانات ----------
    @property
    def today(self):
        return self.d["today"]

    @property
    def meals(self):
        return self.d["today"]["meals"]

    def save(self):
        try:
            tmp = DATA_FILE + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(self.d, f, ensure_ascii=False, indent=2)
            os.replace(tmp, DATA_FILE)
        except Exception as e:
            print("save error:", e)

    def foods(self):
        allf = dict(FOOD_DATABASE)
        allf.update(self.d["custom_foods"])
        return allf

    def get_food(self, name):
        v = self.foods()[name]
        chol = v[6] if len(v) > 6 else CHOL_100G.get(name, 0)
        return tuple(v[:6]) + (chol,)

    def calc_entry(self, name, grams):
        cal, p, c, f, fe, _, chol = self.get_food(name)
        k = grams / 100
        return {"name": f"{name} ({fmt(grams)}ج)", "food": name, "grams": grams,
                "cal": round(cal * k), "p": round(p * k, 1),
                "c": round(c * k, 1), "f": round(f * k, 1), "fe": round(fe * k, 1),
                "chol": round(chol * k)}

    def day_options(self, n=7):
        out = {}
        now = datetime.now()
        for i in range(n):
            d = now - timedelta(days=i)
            name = "اليوم" if i == 0 else "أمس" if i == 1 else f"{DAY_NAMES[d.weekday()]} {d.strftime('%m-%d')}"
            out[name] = d.strftime("%Y-%m-%d")
        return out

    def day_list(self, ds):
        if ds == self.today["date"]:
            return self.today["meals"]
        return self.d["history"].setdefault(ds, {"meals": [], "water": 0.0})["meals"]

    def macro_targets(self):
        g, s = self.d["daily_goal"], self.d["settings"]
        return {"p": g * s.get("p_pct", 25) / 100 / 4,
                "c": g * s.get("c_pct", 50) / 100 / 4,
                "f": g * s.get("f_pct", 25) / 100 / 9}

    def vitc_tip(self, lst, e):
        food = e.get("food")
        if not food or e.get("fe", 0) < 2 or food in VITC_FOODS:
            return ""
        for m in lst:
            if m.get("type") == e.get("type") and m.get("food") in VITC_FOODS:
                return ""
        return "💡 نصيحة: أضف مصدر فيتامين C (برتقال، فلفل رومي، ليمون) مع هالوجبة لزيادة امتصاص الحديد."

    def make_backup(self, auto=False):
        try:
            os.makedirs(BACKUP_DIR, exist_ok=True)
            name = f"auto_{today_str()}.json" if auto else f"backup_{self.stamp()}.json"
            path = os.path.join(BACKUP_DIR, name)
            if auto and os.path.exists(path):
                return None
            with open(path, "w", encoding="utf-8") as f:
                json.dump(self.d, f, ensure_ascii=False, indent=2)
            if auto:
                olds = sorted(x for x in os.listdir(BACKUP_DIR) if x.startswith("auto_"))
                for old in olds[:-7]:
                    os.remove(os.path.join(BACKUP_DIR, old))
            return path
        except Exception as ex:
            print("backup error:", ex)
            return None

    def restore_backup(self, parent):
        path = filedialog.askopenfilename(
            parent=parent, title="اختر ملف النسخة الاحتياطية",
            initialdir=BACKUP_DIR if os.path.isdir(BACKUP_DIR) else BASE_DIR,
            filetypes=[("JSON", "*.json")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                raw = json.load(f)
            if not isinstance(raw, dict) or not ("today" in raw or "daily_goal" in raw):
                raise ValueError
            new = parse_data(raw)
        except Exception:
            messagebox.showerror("خطأ", "الملف مش نسخة صالحة من CalorieMate.", parent=parent)
            return
        if not messagebox.askyesno("استعادة", "رح تُستبدل بياناتك الحالية بالنسخة المختارة.\n"
                                              "(بنحفظ نسخة من بياناتك الحالية أولاً). متأكد؟", parent=parent):
            return
        self.make_backup()
        self.d = new
        roll_day(self.d)
        C.update(THEMES.get(self.d.get("theme", "dark"), THEMES["dark"]))
        self.save()
        parent.destroy()
        self.build_main()
        messagebox.showinfo("تم", "تمت استعادة النسخة الاحتياطية ✅")

    def period_stats(self, offset):
        now = datetime.now()
        logged, water = [], []
        for i in range(7):
            dd = self.day_data((now - timedelta(days=offset + i)).strftime("%Y-%m-%d"))
            if dd["meals"]:
                logged.append(totals(dd["meals"]))
            if dd.get("water", 0) > 0:
                water.append(dd["water"])
        n = len(logged)
        goal = self.d["daily_goal"]

        def avg(k):
            return sum(t[k] for t in logged) / n if n else 0
        return {"days": n, "cal": avg("cal"), "p": avg("p"), "fe": avg("fe"), "chol": avg("chol"),
                "water": sum(water) / len(water) if water else 0,
                "hit": sum(1 for t in logged if goal > 0 and 0.9 * goal <= t["cal"] <= 1.1 * goal)}

    def meal_grams(self, m):
        if m.get("grams"):
            return float(m["grams"])
        mt = re.search(r"\((\d+(?:\.\d+)?)ج\)\s*$", m["name"])
        return float(mt.group(1)) if mt else None

    def rescale_meal(self, m, new_g):
        """تغيير كمية وجبة مسجلة. يعيد الحساب من القاعدة إن وُجد الصنف، وإلا يقيس نسبياً."""
        old_g = self.meal_grams(m)
        if not old_g:
            return False
        base = re.sub(r"\s*\([\d.]+ج\)\s*$", "", m["name"])
        food = m.get("food") or base
        if food in self.foods():
            m.update(self.calc_entry(food, new_g))
            return True
        k = new_g / old_g
        for key, dec in (("cal", 0), ("p", 1), ("c", 1), ("f", 1), ("fe", 1), ("chol", 0)):
            v = m.get(key, 0) * k
            m[key] = round(v) if dec == 0 else round(v, 1)
        m["grams"] = new_g
        m["name"] = f"{base} ({new_g:.0f}ج)"
        return True

    def iron_need(self):
        p = self.d["user_profile"]
        if p.get("gender") == "ذكر":
            return 8
        try:
            age = float(p.get("age", 30))
        except ValueError:
            age = 30
        return 18 if 19 <= age <= 50 else 8

    def chol_limit(self):
        return 200 if "مرتفع" in self.d["user_profile"].get("cholesterol", "") else 300

    def day_data(self, ds):
        if ds == self.today["date"]:
            return self.today
        return self.d["history"].get(ds, {"meals": [], "water": 0.0})

    def new_window(self, title, w, h, modal=False):
        win = tk.Toplevel(self.root)
        win.title(title)
        set_icon(win)
        win.geometry(f"{w}x{h}")
        win.configure(bg=C["MAIN_BG"])
        win.transient(self.root)
        if modal:
            try:
                win.wait_visibility()
                win.grab_set()
            except tk.TclError:
                pass
        card = tk.Frame(win, bg=C["CARD_BG"], padx=15, pady=15)
        card.pack(fill="both", expand=True, padx=12, pady=12)
        return win, card

    # ---------- النافذة الرئيسية ----------
    def setup_styles(self):
        s = ttk.Style()
        s.theme_use("default")
        for name, col in (("Cal", C["ACCENT"]), ("Ok", C["GOOD"]), ("Warn", C["GOLD"]), ("Bad", C["DANGER"]),
                          ("Water", C["WATER"]), ("P", C["BLUE"]), ("C", C["GOLD"]), ("F", C["DANGER"]),
                          ("Fe", C["IRON"])):
            s.configure(f"{name}.Horizontal.TProgressbar", troughcolor=C["CARD_BG"], background=col,
                        lightcolor=col, darkcolor=col, thickness=8, borderwidth=0)

    def build_main(self):
        for w in self.root.winfo_children():
            w.destroy()
        self.root.configure(bg=C["MAIN_BG"])
        self.setup_styles()
        card = tk.Frame(self.root, bg=C["CARD_BG"], padx=18, pady=12)
        card.pack(fill="both", expand=True, padx=16, pady=14)

        header = tk.Frame(card, bg=C["CARD_BG"])
        header.pack(fill="x")
        lbl(header, "CalorieMate Pro 🍎", 20, True).pack()
        btn(header, "🌓", self.toggle_theme, C["MAIN_BG"], fg=C["TEXT"], size=11).place(x=0, y=0)
        lbl(card, "سعراتك، ماؤك، حديدك وكوليسترولك في مكان واحد", 9, fg=C["MUTED"]).pack(pady=(0, 8))

        gf = tk.Frame(card, bg=C["MAIN_BG"], padx=12, pady=6)
        gf.pack(fill="x", pady=(0, 8))
        lbl(gf, "🎯 هدفك اليومي (سعرة):", 10, True, C["GOLD"], bg=C["MAIN_BG"]).pack(side="right")
        self.entry_goal = ent(gf, justify="center", bg=C["CARD_BG"], width=8)
        self.entry_goal.insert(0, str(int(self.d["daily_goal"])))
        self.entry_goal.pack(side="right", padx=8, ipady=2)
        self.entry_goal.bind("<KeyRelease>", self.on_goal_change)

        hero = tk.Frame(card, bg=C["MAIN_BG"], padx=12, pady=8)
        hero.pack(fill="x", pady=(0, 8))
        self.lbl_cal = lbl(hero, "", 22, True, C["GOLD"], bg=C["MAIN_BG"])
        self.lbl_cal.pack()
        self.pb_cal = ttk.Progressbar(hero, style="Cal.Horizontal.TProgressbar", mode="determinate")
        self.pb_cal.pack(fill="x", pady=(4, 4))
        self.lbl_types = lbl(hero, "", 8, True, C["MUTED"], bg=C["MAIN_BG"])
        self.lbl_types.pack()

        mf = tk.Frame(card, bg=C["MAIN_BG"], padx=8, pady=8)
        mf.pack(fill="x", pady=(0, 8))
        self.macro_w = {}
        for col, (key, color, style) in enumerate((("fe", C["IRON"], "Fe"), ("f", C["DANGER"], "F"),
                                                   ("c", C["GOLD"], "C"), ("p", C["BLUE"], "P"))):
            mf.grid_columnconfigure(col, weight=1, uniform="m")
            cell = tk.Frame(mf, bg=C["MAIN_BG"])
            cell.grid(row=0, column=col, padx=5, sticky="ew")
            l = lbl(cell, "", 8, True, color, bg=C["MAIN_BG"], justify="center")
            l.pack()
            pb = ttk.Progressbar(cell, style=f"{style}.Horizontal.TProgressbar", mode="determinate")
            pb.pack(fill="x", pady=(3, 0))
            self.macro_w[key] = (l, pb)

        row2 = tk.Frame(card, bg=C["CARD_BG"])
        row2.pack(fill="x", pady=(0, 6))
        row2.grid_columnconfigure(0, weight=1, uniform="r")
        row2.grid_columnconfigure(1, weight=1, uniform="r")
        cf = tk.Frame(row2, bg=C["MAIN_BG"], padx=10, pady=8)
        cf.grid(row=0, column=1, sticky="nsew", padx=(4, 0))
        self.lbl_chol = lbl(cf, "", 9, True, C["GOOD"], bg=C["MAIN_BG"])
        self.lbl_chol.pack()
        self.pb_chol = ttk.Progressbar(cf, style="Ok.Horizontal.TProgressbar", mode="determinate")
        self.pb_chol.pack(fill="x", pady=(5, 0))
        wf = tk.Frame(row2, bg=C["MAIN_BG"], padx=10, pady=8)
        wf.grid(row=0, column=0, sticky="nsew", padx=(0, 4))
        self.lbl_water = lbl(wf, "", 9, True, C["WATER"], bg=C["MAIN_BG"])
        self.lbl_water.pack()
        self.pb_water = ttk.Progressbar(wf, style="Water.Horizontal.TProgressbar", mode="determinate")
        self.pb_water.pack(fill="x", pady=(5, 5))
        wb = tk.Frame(wf, bg=C["MAIN_BG"])
        wb.pack()
        btn(wb, "+ 250 مل", lambda: self.add_water(0.25), C["WATER"], fg=C["MAIN_BG"], size=8).pack(side="left", padx=2, ipadx=4)
        btn(wb, "- 250", lambda: self.add_water(-0.25), C["CARD_BG"], fg=C["MUTED"], size=8, bold=False).pack(side="left", padx=2, ipadx=4)

        self.lbl_supp = lbl(card, "", 9, True, C["MUTED"])
        self.lbl_supp.pack()
        self.lbl_banner = lbl(card, "", 9, True, C["GOLD"], wraplength=480, justify="center", height=2)
        self.lbl_banner.pack(fill="x")

        btn(card, "➕  إضافة وجبة", self.open_add_meal, C["ACCENT"], size=11).pack(fill="x", ipady=8, pady=(0, 6))

        nav = tk.Frame(card, bg=C["CARD_BG"])
        nav.pack(fill="both", expand=True)
        for c in range(4):
            nav.grid_columnconfigure(c, weight=1, uniform="n")
        for r in range(3):
            nav.grid_rowconfigure(r, weight=1)
        items = [("📋", "وجبات اليوم", self.show_meals), ("⭐", "المفضلة", self.open_favorites),
                 ("🍲", "أكلاتي", self.open_food_manager), ("💊", "المكملات", self.open_supplements),
                 ("📊", "السجل", self.open_history), ("📈", "الملخص", self.open_summary),
                 ("⚖️", "الوزن", self.open_weight), ("🧮", "الاحتياج", self.open_bmr),
                 ("🥦", "أغذية الحديد", self.open_healthy_info), ("📅", "خطة الأسبوع", self.open_plan),
                 ("📤", "تصدير", self.open_export), ("⚙️", "الإعدادات", self.open_settings)]
        for i, (ic, name, cmd) in enumerate(items):
            btn(nav, f"{ic}\n{name}", cmd, C["MAIN_BG"], fg=C["TEXT"], size=8, bold=False).grid(
                row=i // 4, column=3 - (i % 4), sticky="nsew", padx=3, pady=3, ipady=4)
        self.update_ui()

    def toggle_theme(self):
        new = "light" if self.d.get("theme", "dark") == "dark" else "dark"
        self.d["theme"] = new
        C.update(THEMES[new])
        self.save()
        self.build_main()

    def on_goal_change(self, event=None):
        try:
            val = float(self.entry_goal.get().strip())
            if val > 0:
                self.d["daily_goal"] = val
                self.save()
                self.update_ui()
        except ValueError:
            pass

    def set_goal(self, val):
        self.d["daily_goal"] = val
        self.entry_goal.delete(0, tk.END)
        self.entry_goal.insert(0, str(int(val)))
        self.save()
        self.update_ui()

    def show_banner(self, text, color=None, ms=9000):
        if self.banner_job:
            try:
                self.root.after_cancel(self.banner_job)
            except Exception:
                pass
        self.lbl_banner.config(text=text, fg=color or C["GOLD"])
        self.banner_job = self.root.after(ms, lambda: self.lbl_banner.config(text=""))

    def update_ui(self):
        t = totals(self.meals)
        goal = self.d["daily_goal"]
        self.lbl_cal.config(text=f"{round(t['cal'])} / {int(goal)} kcal")
        self.pb_cal["value"] = min(t["cal"] / goal * 100, 100) if goal > 0 else 0
        self.pb_cal.configure(style="Bad.Horizontal.TProgressbar" if t["cal"] > goal
                              else "Cal.Horizontal.TProgressbar")
        by = {k: 0 for k in MEAL_TYPES}
        for m in self.meals:
            by[m.get("type") if m.get("type") in by else "سناك"] += m["cal"]
        self.lbl_types.config(text="   |   ".join(f"{TYPE_ICONS[k]} {k}: {round(v)}" for k, v in by.items()))

        mt = self.macro_targets()
        mt["fe"] = self.iron_need()
        names = {"p": "بروتين", "c": "كارب", "f": "دهون", "fe": "حديد"}
        for key, (l, pb) in self.macro_w.items():
            unit = "mg" if key == "fe" else "ج"
            l.config(text=f"{names[key]}\n{fmt(t[key])} / {fmt(mt[key])}{unit}")
            pb["value"] = min(t[key] / mt[key] * 100, 100) if mt[key] > 0 else 0

        limit = self.chol_limit()
        pct = t["chol"] / limit * 100
        style, col = ("Ok", C["GOOD"]) if pct < 80 else (("Warn", C["GOLD"]) if pct < 100 else ("Bad", C["DANGER"]))
        self.pb_chol["value"] = min(pct, 100)
        self.pb_chol.configure(style=f"{style}.Horizontal.TProgressbar")
        self.lbl_chol.config(text=f"🫀 الكوليسترول\n{round(t['chol'])} / {limit} mg", fg=col)

        water, wgoal = self.today["water"], self.d["water_goal"]
        self.lbl_water.config(text=f"💧 الماء\n{water:.2f} / {wgoal} لتر")
        self.pb_water["value"] = min(water / wgoal * 100, 100) if wgoal > 0 else 0

        supps = self.d["supplements"]
        if supps:
            taken = sum(1 for x in supps if x["name"] in self.today.get("supp_taken", []))
            self.lbl_supp.config(text=f"💊 المكملات اليوم: {taken}/{len(supps)}",
                                 fg=C["GOOD"] if taken == len(supps) else C["MUTED"])
        else:
            self.lbl_supp.config(text="")

        if t["fe"] >= self.iron_need() and not self.today.get("iron_notified"):
            self.today["iron_notified"] = True
            self.save()
            self.show_banner("🎉 ممتاز! وصلت لاحتياجك اليومي من الحديد.", C["IRON"])
        if t["chol"] > limit and not self.today.get("chol_warned"):
            self.today["chol_warned"] = True
            self.save()
            self.show_banner("⚠️ تجاوزت الحد اليومي للكوليسترول، خفف الأكلات الحيوانية بباقي اليوم.", C["DANGER"])

    def add_water(self, amount):
        before = self.today["water"]
        self.today["water"] = round(max(0.0, before + amount), 2)
        if amount > 0:
            self.last_water = datetime.now()
        self.save()
        self.update_ui()
        if amount > 0 and before < self.d["water_goal"] <= self.today["water"]:
            self.show_banner("💧 أحسنت! وصلت لهدفك من الماء اليوم.", C["WATER"])

    def tick(self):
        now = datetime.now()
        if roll_day(self.d):
            self.save()
            self.last_water = now
            self.supp_reminded = set()
            self.update_ui()
            self.show_banner("☀️ يوم جديد! تم حفظ يوم أمس في السجل.", C["GOOD"])
        else:
            try:
                hrs = float(self.d["settings"].get("water_reminder_hrs", 2) or 0)
            except (TypeError, ValueError):
                hrs = 2.0
            hm = now.strftime("%H:%M")
            taken = self.today.get("supp_taken", [])
            due = [x for x in self.d["supplements"]
                   if x.get("time") and hm >= x["time"] and x["name"] not in taken
                   and x["name"] not in self.supp_reminded]
            if due:
                self.supp_reminded.add(due[0]["name"])
                self.show_banner(f"💊 وقت {due[0]['name']}", C["GOLD"], 20000)
            elif hrs > 0 and 7 <= now.hour < 23 and self.today["water"] < self.d["water_goal"] \
                    and now - self.last_water >= timedelta(hours=hrs):
                self.last_water = now
                self.show_banner("💧 وقت الماء! اشرب كوب.", C["WATER"], 15000)
        self.root.after(60000, self.tick)

    def reset_day(self):
        if messagebox.askyesno("تصفير اليوم", "تأكد: بدك تمسح وجبات اليوم والماء وتبدأ من جديد؟\n(ما بينحفظ بالسجل)"):
            self.d["today"] = new_today()
            self.save()
            self.update_ui()

    # ---------- إضافة وجبة ----------
    def open_add_meal(self):
        win, card = self.new_window("إضافة وجبة 🍽️", 520, 780, True)
        lbl(card, "إضافة وجبة", 14, True).pack(pady=(0, 6))
        opts = self.day_options()
        day_var = tk.StringVar(value="اليوم")
        hr = datetime.now().hour
        meal_type = tk.StringVar(value="فطور" if hr < 11 else "غداء" if hr < 16 else "عشاء" if hr < 21 else "سناك")
        top = tk.Frame(card, bg=C["CARD_BG"])
        top.pack(fill="x", pady=(0, 6))
        lbl(top, "النوع / اليوم:", 9, fg=C["MUTED"]).pack(side="right")
        optmenu(top, meal_type, *MEAL_TYPES).pack(side="right", padx=6)
        optmenu(top, day_var, *opts.keys()).pack(side="right")

        iron_only = tk.BooleanVar(value=False)
        holder = {"entry": None}
        state = {"food": None, "units": {"جرام": 1.0}}

        def refresh(*_):
            food_listbox.delete(0, tk.END)
            q = norm(search_entry.get())
            usage = self.d.setdefault("usage", {})
            for item in sorted(self.foods(), key=lambda n: (-usage.get(n, 0), n)):
                if iron_only.get() and self.get_food(item)[4] < 2.0:
                    continue
                if q in norm(item):
                    food_listbox.insert(tk.END, item)
            preview()

        tk.Checkbutton(card, text="🩸 الأكلات الغنية بالحديد فقط (أعلى من 2mg)", variable=iron_only,
                       bg=C["CARD_BG"], fg=C["IRON"], selectcolor=C["MAIN_BG"], activebackground=C["CARD_BG"],
                       activeforeground=C["IRON"], font=(FONT, 8, "bold"), command=refresh).pack(anchor="e", pady=(0, 4))
        lbl(card, "🔍 ابحث عن الأكلة (Enter لاختيار أول نتيجة):", 9, True, C["GOLD"]).pack(anchor="e")
        search_entry = ent(card)
        search_entry.pack(fill="x", pady=(2, 6), ipady=3)
        food_listbox = tk.Listbox(card, font=(FONT, 9), bg=C["MAIN_BG"], fg=C["TEXT"],
                                  selectbackground=C["ACCENT"], selectforeground="white", bd=0,
                                  highlightthickness=0, height=6, justify="right", exportselection=False)
        food_listbox.pack(fill="x", pady=(0, 6))
        search_entry.bind("<KeyRelease>", refresh)

        qf = tk.Frame(card, bg=C["CARD_BG"])
        qf.pack(fill="x", pady=(2, 6))
        lbl(qf, "⚖️ الكمية والوحدة:", 9, True, C["GOLD"]).pack(side="right")
        entry_qty = ent(qf, justify="center", width=7)
        entry_qty.insert(0, "100")
        entry_qty.pack(side="right", padx=8, ipady=3)
        unit_var = tk.StringVar(value="جرام")
        unit_menu = optmenu(qf, unit_var, "جرام")
        unit_menu.pack(side="right")
        btn(qf, "💡 حصة تقريبية", lambda: suggest(), C["CARD_BG"], fg=C["GOLD"], size=8).pack(side="left", padx=(0, 6))
        lbl_grams = lbl(qf, "", 8, fg=C["MUTED"])
        lbl_grams.pack(side="left")

        pf = tk.Frame(card, bg=C["MAIN_BG"], padx=10, pady=8)
        pf.pack(fill="x", pady=(0, 6))
        lbl_res = lbl(pf, "اختر أكلة للمعاينة...", 11, True, C["GOLD"], bg=C["MAIN_BG"])
        lbl_res.pack()
        lbl_mac = lbl(pf, "الماكروز: -", 8, fg=C["MUTED"], bg=C["MAIN_BG"])
        lbl_mac.pack(pady=(3, 0))

        def units_for(food):
            portion = float(self.get_food(food)[5])
            u = {"جرام": 1.0, f"حصة ({fmt(portion)}ج)": portion}
            for k, v in MEASURES.get(food, {}).items():
                u[f"{k} ({fmt(v)}ج)"] = float(v)
            u["ملعقة صغيرة (5ج)"] = 5.0
            u["ملعقة كبيرة (15ج)"] = 15.0
            return u

        def set_units(food):
            was_gram = unit_var.get() == "جرام"
            u = units_for(food)
            state["units"], state["food"] = u, food
            menu = unit_menu["menu"]
            menu.delete(0, "end")
            for name in u:
                menu.add_command(label=name, command=lambda v=name: (unit_var.set(v), preview()))
            unit_var.set("جرام")  # الافتراضي جرام، والوحدات الأخرى اختيارية
            if not was_gram or not entry_qty.get().strip():
                entry_qty.delete(0, tk.END)
                entry_qty.insert(0, "100")

        def suggest():
            sel = food_listbox.curselection()
            if not sel:
                messagebox.showinfo("تنبيه", "اختر أكلة أولاً لاقتراح حصتها.", parent=win)
                return
            portion = float(self.get_food(food_listbox.get(sel[0]))[5])
            unit_var.set("جرام")
            entry_qty.delete(0, tk.END)
            entry_qty.insert(0, fmt(portion))
            preview()

        def preview(*_):
            sel = food_listbox.curselection()
            if not sel:
                holder["entry"] = None
                lbl_res.config(text="اختر أكلة من القائمة")
                lbl_mac.config(text="الماكروز: -")
                lbl_grams.config(text="")
                return
            name = food_listbox.get(sel[0])
            if name != state["food"]:
                set_units(name)
            try:
                qty = float(entry_qty.get().strip().replace(",", "."))
                if qty <= 0:
                    raise ValueError
            except ValueError:
                holder["entry"] = None
                lbl_res.config(text="أدخل كمية صحيحة")
                lbl_mac.config(text="الماكروز: -")
                return
            unit = unit_var.get()
            grams = qty * state["units"].get(unit, 1.0)
            e = self.calc_entry(name, grams)
            if unit != "جرام":
                e["measure"] = f"{fmt(qty)} {unit.split(' (')[0]}"
            holder["entry"] = e
            lbl_grams.config(text=f"= {fmt(grams)} جرام")
            lbl_res.config(text=f"النتيجة: {e['cal']} سعرة حرارية")
            lbl_mac.config(text=f"بروتين: {e['p']}ج | كارب: {e['c']}ج | دهون: {e['f']}ج | "
                                f"حديد: {e['fe']}mg | كوليسترول: {e['chol']}mg")

        def pick_first(event=None):
            if food_listbox.size():
                food_listbox.selection_clear(0, tk.END)
                food_listbox.selection_set(0)
                preview()
                entry_qty.focus_set()
                entry_qty.select_range(0, tk.END)

        food_listbox.bind("<<ListboxSelect>>", preview)
        entry_qty.bind("<KeyRelease>", preview)
        entry_qty.bind("<FocusIn>", lambda e: entry_qty.select_range(0, tk.END))
        search_entry.bind("<Return>", pick_first)
        entry_qty.bind("<Return>", lambda e: add_to_day())

        lbl(card, "📝 ملاحظة (اختياري):", 9, fg=C["MUTED"]).pack(anchor="e")
        note_entry = ent(card)
        note_entry.pack(fill="x", pady=(2, 6), ipady=3)
        lbl_msg = lbl(card, "", 9, True, C["GOOD"], wraplength=470, justify="center")
        lbl_msg.pack(pady=(0, 6))

        def add_to_day():
            e = holder["entry"]
            if not e:
                messagebox.showwarning("تنبيه", "اختر أكلة وحدد كمية صحيحة أولاً!", parent=win)
                return
            ds = opts[day_var.get()]
            e = dict(e, type=meal_type.get())
            note = note_entry.get().strip()
            if note:
                e["note"] = note
            lst = self.day_list(ds)
            lst.append(e)
            usage = self.d.setdefault("usage", {})
            usage[e["food"]] = usage.get(e["food"], 0) + 1
            self.save()
            self.update_ui()
            msg = f"✓ تمت إضافة ({e['food']}) إلى {meal_type.get()} — {day_var.get()}"
            if self.d["settings"].get("vitc_tip", True):
                tip = self.vitc_tip(lst, e)
                if tip:
                    msg += "\n" + tip
                    if ds == self.today["date"]:
                        self.show_banner(tip, C["BLUE"], 12000)
            lbl_msg.config(text=msg)
            note_entry.delete(0, tk.END)
            search_entry.delete(0, tk.END)
            state["food"] = None
            refresh()
            search_entry.focus_set()

        btn(card, "➕ إضافة هذه الوجبة ومتابعة الإضافة", add_to_day, C["ACCENT"], size=10).pack(fill="x", ipady=7)
        refresh()
        search_entry.focus_set()

    # ---------- وجبات اليوم + حفظ كمفضلة ----------
    def show_meals(self):
        win, card = self.new_window("وجبات اليوم 📋", 620, 740)
        opts = self.day_options()
        day_var = tk.StringVar(value="اليوم")
        hdr = tk.Frame(card, bg=C["CARD_BG"])
        hdr.pack(fill="x", pady=(0, 4))
        lbl(hdr, "قائمة الوجبات", 14, True).pack(side="right")
        optmenu(hdr, day_var, *opts.keys(), command=lambda v: render()).pack(side="left")
        lbl(card, "غيّر نوع الوجبة من القائمة المنسدلة، عدّل الكمية، أضف ملاحظة، أو كرّر الوجبة", 8,
            fg=C["MUTED"]).pack(pady=(0, 6))
        bottom = tk.Frame(card, bg=C["CARD_BG"])
        bottom.pack(side="bottom", fill="x", pady=(8, 0))
        lbl_total = lbl(card, "", 9, True, C["GOLD"])
        lbl_total.pack(side="bottom", pady=(6, 0))
        holder = tk.Frame(card, bg=C["CARD_BG"])
        holder.pack(fill="both", expand=True)
        inner = make_scrollable(holder)
        undo = {"ds": None, "idx": 0, "m": None}

        def lst():
            return self.day_list(opts[day_var.get()])

        def refresh():
            self.save()
            self.update_ui()
            render()

        def index_of(L, m):
            return next(i for i, x in enumerate(L) if x is m)

        def set_type(m, new):
            m["type"] = new
            refresh()

        def delete(m):
            L = lst()
            i = index_of(L, m)
            L.pop(i)
            undo.update(ds=opts[day_var.get()], idx=i, m=m)
            btn_undo.config(state="normal", text=f"↩ تراجع عن حذف ({(m.get('food') or m['name'])[:16]})")
            refresh()

        def do_undo():
            if undo["m"] is None:
                return
            L = self.day_list(undo["ds"])
            L.insert(min(undo["idx"], len(L)), undo["m"])
            undo["m"] = None
            btn_undo.config(state="disabled", text="↩ تراجع")
            refresh()

        def duplicate(m):
            L = lst()
            L.insert(index_of(L, m) + 1, dict(m))
            refresh()

        def edit_note(m):
            v = simpledialog.askstring("ملاحظة", "ملاحظة على هالوجبة (اتركها فاضية للحذف):",
                                       initialvalue=m.get("note", ""), parent=win)
            if v is None:
                return
            if v.strip():
                m["note"] = v.strip()
            else:
                m.pop("note", None)
            refresh()

        def edit_grams(m):
            cur = self.meal_grams(m)
            if not cur:
                messagebox.showinfo("تنبيه", "هاي وجبة قديمة وما فيها كمية محفوظة، احذفها وأضفها من جديد.", parent=win)
                return
            g = simpledialog.askfloat("تعديل الكمية", "الكمية الجديدة (جرام):", initialvalue=cur,
                                      minvalue=1, maxvalue=5000, parent=win)
            if g:
                m.pop("measure", None)
                self.rescale_meal(m, g)
                refresh()

        def render():
            for w in inner.winfo_children():
                w.destroy()
            L = lst()
            t = totals(L)
            lbl_total.config(text=f"المجموع: {round(t['cal'])} kcal | 🩸 {round(t['fe'], 1)}mg | 🫀 {round(t['chol'])}mg")
            if not L:
                lbl(inner, "لا توجد وجبات مسجلة بهاليوم.", 11, fg=C["MUTED"]).pack(pady=40)
                return
            groups = {k: [] for k in MEAL_TYPES}
            for m in L:
                groups[m.get("type") if m.get("type") in groups else "سناك"].append(m)
            for mt in MEAL_TYPES:
                items = groups[mt]
                if not items:
                    continue
                lbl(inner, f"{TYPE_ICONS[mt]} {mt}  —  {round(sum(x['cal'] for x in items))} kcal",
                    10, True, C["GOLD"]).pack(anchor="e", pady=(10, 2))
                for m in items:
                    row = tk.Frame(inner, bg=C["MAIN_BG"], padx=10, pady=6)
                    row.pack(fill="x", pady=3)
                    top = tk.Frame(row, bg=C["MAIN_BG"])
                    top.pack(fill="x")
                    lbl(top, m["name"], 9, True, bg=C["MAIN_BG"]).pack(side="right")
                    lbl(top, f"{m['cal']} kcal", 9, True, C["GOLD"], bg=C["MAIN_BG"]).pack(side="left")
                    extra = f"   • {m['measure']}" if m.get("measure") else ""
                    lbl(row, f"🩸 {m.get('fe', 0)}mg   🫀 {m.get('chol', 0)}mg   بروتين {m.get('p', 0)}ج{extra}",
                        8, fg=C["MUTED"], bg=C["MAIN_BG"]).pack(anchor="e")
                    if m.get("note"):
                        lbl(row, f"📝 {m['note']}", 8, fg=C["GOLD"], bg=C["MAIN_BG"],
                            wraplength=540, justify="right").pack(anchor="e")
                    sub = tk.Frame(row, bg=C["MAIN_BG"])
                    sub.pack(fill="x", pady=(4, 0))
                    btn(sub, "🗑", lambda m=m: delete(m), C["DANGER"]).pack(side="left", padx=2)
                    btn(sub, "⧉", lambda m=m: duplicate(m), C["ACCENT"]).pack(side="left", padx=2)
                    btn(sub, "📝", lambda m=m: edit_note(m), C["CARD_BG"], fg=C["TEXT"]).pack(side="left", padx=2)
                    btn(sub, "✏️ الكمية", lambda m=m: edit_grams(m), C["GOOD"], size=8).pack(side="left", padx=2)
                    tv = tk.StringVar(value=mt)
                    optmenu(sub, tv, *MEAL_TYPES,
                            command=lambda v, m=m: win.after(10, lambda: set_type(m, v))).pack(side="left", padx=4)

        filt = tk.StringVar(value="الكل")

        def save_fav():
            items = [m for m in lst() if filt.get() == "الكل" or m.get("type") == filt.get()]
            if not items:
                messagebox.showwarning("تنبيه", "ما في وجبات بهاد النوع لحفظها.", parent=win)
                return
            name = simpledialog.askstring("مفضلة", "اسم الوجبة المفضلة (مثلاً: فطوري المعتاد):", parent=win)
            if name and name.strip():
                self.d["favorites"][name.strip()] = [{k: v for k, v in m.items() if k not in ("type", "note")}
                                                     for m in items]
                self.save()
                messagebox.showinfo("تم", f"تم حفظ ({name.strip()}) بالمفضلة ⭐", parent=win)

        btn(bottom, "⭐ حفظ كمفضلة", save_fav, C["GOLD"], fg="black").pack(side="right", ipadx=8, ipady=3)
        optmenu(bottom, filt, "الكل", *MEAL_TYPES).pack(side="right", padx=6)
        lbl(bottom, "حفظ:", 9, fg=C["MUTED"]).pack(side="right")
        btn_undo = btn(bottom, "↩ تراجع", do_undo, C["MAIN_BG"], fg=C["TEXT"], state="disabled")
        btn_undo.pack(side="left", ipadx=6, ipady=3)
        render()

    # ---------- المفضلة ----------
    def open_favorites(self):
        win, card = self.new_window("وجباتي المفضلة ⭐", 470, 520)
        lbl(card, "وجباتي المفضلة", 13, True, C["GOLD"]).pack(pady=(0, 6))
        tf = tk.Frame(card, bg=C["CARD_BG"])
        tf.pack(fill="x", pady=(0, 8))
        lbl(tf, "أضفها كـ:", 9, fg=C["MUTED"]).pack(side="right")
        mtype = tk.StringVar(value="فطور")
        optmenu(tf, mtype, *MEAL_TYPES).pack(side="right", padx=6)
        holder = tk.Frame(card, bg=C["CARD_BG"])
        holder.pack(fill="both", expand=True)
        inner = make_scrollable(holder)

        def add_fav(name):
            for m in self.d["favorites"][name]:
                self.meals.append(dict(m, type=mtype.get()))
            self.save()
            self.update_ui()
            self.show_banner(f"⭐ تمت إضافة ({name}) إلى {mtype.get()}", C["GOOD"])

        def del_fav(name):
            if messagebox.askyesno("حذف", f"حذف ({name}) من المفضلة؟", parent=win):
                self.d["favorites"].pop(name, None)
                self.save()
                render()

        def render():
            for w in inner.winfo_children():
                w.destroy()
            favs = self.d["favorites"]
            if not favs:
                lbl(inner, "ما في مفضلة بعد.\nافتح «وجبات اليوم» واضغط «حفظ كمفضلة».", 10,
                    fg=C["MUTED"], justify="center").pack(pady=40)
                return
            for name in sorted(favs):
                items = favs[name]
                row = tk.Frame(inner, bg=C["MAIN_BG"], padx=10, pady=6)
                row.pack(fill="x", pady=3)
                lbl(row, name, 10, True, bg=C["MAIN_BG"]).pack(side="right")
                lbl(row, f"{len(items)} أصناف | {round(sum(m['cal'] for m in items))} kcal", 8,
                    fg=C["GOLD"], bg=C["MAIN_BG"]).pack(side="right", padx=8)
                btn(row, "🗑", lambda n=name: del_fav(n), C["DANGER"]).pack(side="left", padx=2)
                btn(row, "➕ أضف", lambda n=name: add_fav(n), C["GOOD"]).pack(side="left", padx=2)
        render()

    # ---------- الأكلات المخصصة ----------
    def open_food_manager(self):
        win, card = self.new_window("أكلاتي المخصصة 🍲", 460, 540)
        lbl(card, "الأكلات التي أضفتها (لكل 100 جرام)", 12, True, C["GOLD"]).pack(pady=(0, 8))
        bar = tk.Frame(card, bg=C["CARD_BG"])
        bar.pack(fill="x", pady=(0, 8))
        btn(bar, "➕ أكلة جديدة", lambda: self.open_food_form(None, render), C["GOOD"], size=10).pack(
            side="right", expand=True, fill="x", padx=2, ipady=5)
        btn(bar, "🍳 وصفة مركبة", lambda: self.open_recipe(render), C["ACCENT"], size=10).pack(
            side="right", expand=True, fill="x", padx=2, ipady=5)
        holder = tk.Frame(card, bg=C["CARD_BG"])
        holder.pack(fill="both", expand=True)
        inner = make_scrollable(holder)

        def delete(name):
            if messagebox.askyesno("حذف", f"حذف ({name}) نهائياً؟", parent=win):
                self.d["custom_foods"].pop(name, None)
                self.save()
                render()

        def render():
            for w in inner.winfo_children():
                w.destroy()
            foods = self.d["custom_foods"]
            if not foods:
                lbl(inner, "لا توجد أكلات مخصصة بعد.", 10, fg=C["MUTED"]).pack(pady=40)
                return
            for name in sorted(foods):
                v = foods[name]
                row = tk.Frame(inner, bg=C["MAIN_BG"], padx=8, pady=6)
                row.pack(fill="x", pady=3)
                lbl(row, name, 9, True, bg=C["MAIN_BG"]).pack(side="right")
                lbl(row, f"{round(v[0])} kcal", 8, fg=C["GOLD"], bg=C["MAIN_BG"]).pack(side="right", padx=8)
                btn(row, "🗑", lambda n=name: delete(n), C["DANGER"]).pack(side="left", padx=2)
                btn(row, "✏️", lambda n=name: self.open_food_form(n, render), C["GOOD"]).pack(side="left", padx=2)
        render()

    def open_food_form(self, old_name=None, on_done=None):
        win, card = self.new_window("تعديل أكلة ✏️" if old_name else "إضافة أكلة جديدة ➕", 400, 640, True)
        lbl(card, "بيانات الصنف (لكل 100 جرام)", 11, True, C["GOLD"]).pack(pady=(0, 8))
        old = self.d["custom_foods"].get(old_name) if old_name else None
        spec = [("اسم الطعام:", "name"), ("السعرات (kcal):", "cal"), ("البروتين (جرام):", "p"),
                ("الكاربوهيدرات (جرام):", "c"), ("الدهون (جرام):", "f"), ("الحديد (ملغرام):", "fe"),
                ("وزن الحصة التقريبي (جرام):", "portion"), ("الكوليسترول (ملغرام):", "chol")]
        fields = {}
        for i, (text, key) in enumerate(spec):
            lbl(card, text, 8, fg=C["MUTED"]).pack(anchor="e")
            e = ent(card)
            if key == "name":
                e.insert(0, old_name or "")
            elif old is not None:
                val = old[i - 1] if i - 1 < len(old) else 0
                e.insert(0, str(val))
            else:
                e.insert(0, "100" if key == "portion" else ("0" if key == "chol" else ""))
            e.pack(fill="x", pady=(1, 5), ipady=3)
            fields[key] = e

        def save_food():
            name = fields["name"].get().strip()
            if not name:
                messagebox.showerror("خطأ", "اكتب اسم الصنف!", parent=win)
                return
            try:
                nums = [float(fields[k].get()) for k, _ in [(s[1], 0) for s in spec[1:]]]
                if any(x < 0 for x in nums):
                    raise ValueError
            except ValueError:
                messagebox.showerror("خطأ", "أدخل أرقاماً صحيحة (صفر أو أكثر) لكل القيم!", parent=win)
                return
            if nums[5] <= 0:
                nums[5] = 100.0
            foods = self.d["custom_foods"]
            if old_name and old_name != name:
                foods.pop(old_name, None)
            foods[name] = nums  # [cal, p, c, f, fe, portion, chol]
            self.save()
            win.destroy()
            if on_done:
                on_done()

        btn(card, "💾 حفظ", save_food, C["ACCENT"], size=10).pack(fill="x", pady=(8, 0), ipady=6)

    # ---------- حاسبة الاحتياج ----------
    def open_bmr(self):
        win, card = self.new_window("حاسبة الاحتياج اليومي 🧮", 480, 720, True)
        p = self.d["user_profile"]
        lbl(card, "حاسبة الاحتياج اليومي الشاملة", 13, True).pack(pady=(0, 6))
        lbl(card, "الجنس:", 8, fg=C["MUTED"]).pack(anchor="e")
        gender = tk.StringVar(value=p.get("gender", "أنثى"))
        gf = tk.Frame(card, bg=C["CARD_BG"])
        gf.pack(fill="x")
        for text, pad in (("أنثى", 0), ("ذكر", 10)):
            tk.Radiobutton(gf, text=text, variable=gender, value=text, bg=C["CARD_BG"], fg=C["TEXT"],
                           selectcolor=C["MAIN_BG"], activebackground=C["CARD_BG"],
                           activeforeground=C["TEXT"]).pack(side="right", padx=pad)

        def field(text, key, default):
            lbl(card, text, 8, fg=C["MUTED"]).pack(anchor="e")
            e = ent(card)
            e.insert(0, p.get(key, default))
            e.pack(fill="x", pady=1, ipady=2)
            return e
        e_w = field("الوزن (كجم):", "weight", "65")
        e_h = field("الطول (سم):", "height", "165")
        e_a = field("العمر (سنة):", "age", "30")

        lbl(card, "مستوى النشاط البدني:", 8, fg=C["MUTED"]).pack(anchor="e")
        activity = tk.StringVar(value=p.get("activity", "قليل النشاط (مكتبي)"))
        optmenu(card, activity, "خامل (بدون تمارين)", "قليل النشاط (مكتبي)",
                "متوسط (3-5 أيام رياضية)", "نشيط جداً (رياضي يومي)").pack(fill="x", pady=(1, 4))
        lbl(card, "حالة الكوليسترول لديك:", 8, True, C["DANGER"]).pack(anchor="e")
        chol = tk.StringVar(value=p.get("cholesterol", "مرتفع (حد أقصى 200 mg)"))
        optmenu(card, chol, "مرتفع (حد أقصى 200 mg)", "طبيعي (حد أقصى 300 mg)").pack(fill="x", pady=(1, 6))

        rf = tk.Frame(card, bg=C["MAIN_BG"], padx=10, pady=8)
        rf.pack(fill="x", pady=4)
        lbl_main = lbl(rf, "اضغط احسب للحصول على النتيجة...", 8, True, C["MUTED"], bg=C["MAIN_BG"])
        lbl_main.pack()
        lbl_micro = lbl(rf, "", 8, fg=C["GOLD"], bg=C["MAIN_BG"])
        lbl_micro.pack(pady=(2, 0))
        goals = {"lose": 0, "maintain": 0, "gain": 0}

        def pick(key):
            if goals[key] > 0:
                self.set_goal(goals[key])
                messagebox.showinfo("نجاح", f"تم تحديد {int(goals[key])} سعرة كهدف يومي جديد!", parent=win)

        bf = tk.Frame(card, bg=C["CARD_BG"])
        bf.pack(fill="x", pady=(4, 0))
        b_lose = btn(bf, "📉 خسارة وزن", lambda: pick("lose"), C["DANGER"], size=8, state="disabled")
        b_main = btn(bf, "⚖️ ثبات وزن", lambda: pick("maintain"), C["GOOD"], size=8, state="disabled")
        b_gain = btn(bf, "📈 زيادة وزن", lambda: pick("gain"), C["GOLD"], fg="black", size=8, state="disabled")
        for b in (b_lose, b_main, b_gain):
            b.pack(side="left", expand=True, fill="x", padx=2)

        def calculate():
            try:
                w, h, a = float(e_w.get()), float(e_h.get()), float(e_a.get())
                if w <= 0 or h <= 0 or a <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("خطأ", "اكتب أرقاماً صحيحة للوزن والطول والعمر", parent=win)
                return
            p.update({"gender": gender.get(), "weight": str(w), "height": str(h), "age": str(a),
                      "activity": activity.get(), "cholesterol": chol.get()})
            self.save()
            bmr = 10 * w + 6.25 * h - 5 * a + (5 if gender.get() == "ذكر" else -161)
            act = activity.get()
            mult = 1.2
            if "قليل" in act:
                mult = 1.375
            elif "متوسط" in act:
                mult = 1.55
            elif "نشيط" in act:
                mult = 1.725
            tdee = round(bmr * mult)
            goals.update({"maintain": tdee, "lose": max(1200, tdee - 500), "gain": tdee + 500})
            lbl_main.config(text=f"📉 خسارة: {goals['lose']} | ⚖️ ثبات: {tdee} | 📈 زيادة: {goals['gain']}", fg=C["GOLD"])
            lbl_micro.config(text=f"🩸 الاحتياج اليومي للحديد: {self.iron_need()} mg\n"
                                  f"🫀 حد الكوليسترول اليومي: أقل من {self.chol_limit()} mg")
            for b in (b_lose, b_main, b_gain):
                b.config(state="normal")
            self.update_ui()

        btn(card, "احسب واحفظ البيانات", calculate, C["ACCENT"], size=10).pack(fill="x", ipady=5, pady=(8, 5))
        calculate()

    # ---------- سجل الأيام ----------
    def open_history(self):
        win, card = self.new_window("سجل الأيام 📊", 500, 640)
        lbl(card, "آخر 7 أيام", 13, True, C["GOLD"]).pack(pady=(0, 6))
        today = datetime.now()
        dates = [(today - timedelta(days=i)) for i in range(6, -1, -1)]
        ds_list = [d.strftime("%Y-%m-%d") for d in dates]
        labels = [d.strftime("%m-%d") for d in dates]
        tots = [totals(self.day_data(ds)["meals"]) for ds in ds_list]
        waters = [self.day_data(ds).get("water", 0.0) for ds in ds_list]
        metrics = [
            ("السعرات", [t["cal"] for t in tots], self.d["daily_goal"], C["ACCENT"]),
            ("الكوليسترول", [t["chol"] for t in tots], self.chol_limit(), C["IRON"]),
            ("الحديد", [t["fe"] for t in tots], self.iron_need(), C["GOLD"]),
            ("الماء", waters, self.d["water_goal"], C["WATER"]),
        ]
        canvas = tk.Canvas(card, width=450, height=220, bg=C["MAIN_BG"], highlightthickness=0)
        info = lbl(card, "", 9, fg=C["MUTED"])

        def show(i):
            name, vals, limit, col = metrics[i]
            draw_bars(canvas, labels, vals, limit, col, 450, 220)
            used = [v for v in vals if v > 0]
            avg = sum(used) / len(used) if used else 0
            info.config(text=f"المتوسط ({name}) للأيام المسجلة: {fmt(avg)}   |   الخط المتقطع = الهدف/الحد")
            for j, b in enumerate(btns):
                b.config(bg=C["ACCENT"] if j == i else C["MAIN_BG"], fg="white" if j == i else C["TEXT"])

        tabs = tk.Frame(card, bg=C["CARD_BG"])
        tabs.pack(fill="x", pady=(0, 6))
        btns = []
        for i, m in enumerate(metrics):
            b = btn(tabs, m[0], lambda i=i: show(i), C["MAIN_BG"], fg=C["TEXT"], size=8)
            b.pack(side="right", expand=True, fill="x", padx=1)
            btns.append(b)
        canvas.pack(pady=4)
        info.pack(pady=(0, 6))

        table = tk.Frame(card, bg=C["MAIN_BG"], padx=6, pady=6)
        table.pack(fill="x")
        for c, h in enumerate(("اليوم", "سعرات", "حديد", "كوليسترول", "ماء")):
            lbl(table, h, 8, True, C["GOLD"], bg=C["MAIN_BG"]).grid(row=0, column=4 - c, padx=8, sticky="e")
        for r, i in enumerate(reversed(range(7)), start=1):
            dname = DAY_NAMES[dates[i].weekday()]
            row = (f"{labels[i]} {dname}", fmt(tots[i]["cal"]), fmt(tots[i]["fe"]),
                   fmt(tots[i]["chol"]), f"{waters[i]:.2f}")
            for c, val in enumerate(row):
                lbl(table, val, 8, bg=C["MAIN_BG"]).grid(row=r, column=4 - c, padx=8, sticky="e")
        show(0)

    # ---------- الوزن و BMI ----------
    def open_weight(self):
        win, card = self.new_window("الوزن و BMI ⚖️", 480, 700)
        p = self.d["user_profile"]
        lbl(card, "تتبع الوزن", 13, True, C["GOLD"]).pack(pady=(0, 6))
        rf = tk.Frame(card, bg=C["CARD_BG"])
        rf.pack(fill="x", pady=(0, 6))
        lbl(rf, "وزنك اليوم (كجم):", 9, True, C["GOLD"]).pack(side="right")
        e = ent(rf, justify="center", width=8)
        ws = self.d["weights"]
        e.insert(0, str(ws[-1]["weight"]) if ws else p.get("weight", "65"))
        e.pack(side="right", padx=6, ipady=3)
        lbl_bmi = lbl(card, "", 10, True, C["GOOD"])
        lbl_chg = lbl(card, "", 9, fg=C["MUTED"])
        canvas = tk.Canvas(card, width=440, height=200, bg=C["MAIN_BG"], highlightthickness=0)
        holder = tk.Frame(card, bg=C["CARD_BG"])
        inner_parent = holder

        def render():
            ws = self.d["weights"]
            try:
                h = float(p.get("height", 165)) / 100
                cur = ws[-1]["weight"] if ws else float(p.get("weight", 65))
                bmi = cur / (h * h)
                cat = ("نحافة" if bmi < 18.5 else "وزن طبيعي" if bmi < 25
                       else "زيادة وزن" if bmi < 30 else "سمنة")
                lbl_bmi.config(text=f"BMI: {bmi:.1f}  —  {cat}")
            except (ValueError, ZeroDivisionError):
                lbl_bmi.config(text="BMI: أدخل طولك من حاسبة الاحتياج")
            if len(ws) >= 2:
                diff = ws[-1]["weight"] - ws[0]["weight"]
                lbl_chg.config(text=f"التغير منذ أول قياس: {diff:+.1f} كجم")
            else:
                lbl_chg.config(text="")
            last = ws[-12:]
            draw_line(canvas, [x["date"][5:] for x in last], [x["weight"] for x in last], 440, 200)
            for w in list_inner.winfo_children():
                w.destroy()
            for x in reversed(ws[-8:]):
                row = tk.Frame(list_inner, bg=C["MAIN_BG"], padx=8, pady=4)
                row.pack(fill="x", pady=2)
                lbl(row, f"{x['date']}", 9, bg=C["MAIN_BG"]).pack(side="right")
                lbl(row, f"{x['weight']} كجم", 9, True, C["GOLD"], bg=C["MAIN_BG"]).pack(side="right", padx=10)
                btn(row, "🗑", lambda d=x["date"]: delete(d), C["DANGER"]).pack(side="left")

        def delete(date):
            self.d["weights"] = [x for x in self.d["weights"] if x["date"] != date]
            self.save()
            render()

        def save_w():
            try:
                w = float(e.get())
                if not 20 <= w <= 400:
                    raise ValueError
            except ValueError:
                messagebox.showerror("خطأ", "أدخل وزناً صحيحاً (20 - 400 كجم)", parent=win)
                return
            ds = today_str()
            lst = [x for x in self.d["weights"] if x["date"] != ds]
            lst.append({"date": ds, "weight": w})
            lst.sort(key=lambda x: x["date"])
            self.d["weights"] = lst
            p["weight"] = str(w)
            self.save()
            render()

        btn(rf, "💾 سجّل", save_w, C["ACCENT"]).pack(side="left", ipadx=8, ipady=2)
        lbl_bmi.pack()
        lbl_chg.pack(pady=(0, 4))
        canvas.pack(pady=4)
        holder.pack(fill="both", expand=True)
        list_inner = make_scrollable(inner_parent)
        render()

    # ---------- معلومات وخطة ----------
    def open_healthy_info(self):
        win, card = self.new_window("أطعمة ممتازة للحديد والكوليسترول 🥦", 470, 560)
        lbl(card, "أطعمة ترفع الحديد وتخفض الكوليسترول", 12, True, C["GOLD"]).pack(pady=(0, 10))
        txt = tk.Text(card, font=(FONT, 9), bg=C["MAIN_BG"], fg=C["TEXT"], bd=0, padx=10, pady=10, wrap="word")
        txt.pack(fill="both", expand=True)
        txt.insert(tk.END, """
🌿 1. العدس والبقوليات:
• غنية بالحديد وخالية من الكوليسترول.
• فيها ألياف ذائبة تساعد على خفض الكوليسترول الضار.

🥬 2. السبانخ والخضار الورقية:
• مصدر ممتاز للحديد والألياف.
• أضف الليمون (فيتامين C) لزيادة امتصاص الحديد.

🥣 3. الشوفان وبذور الكتان:
• الشوفان فيه بيتا-جلوكان لخفض الكوليسترول.
• بذور الكتان فيها أوميغا-3.

🍗 4. صدور الدجاج والأسماك المشوية:
• بروتين صافي وكوليسترول أقل من اللحوم الحمراء.

⚠️ 5. الكبدة والجمبري والبيض:
• الكبدة غنية جداً بالحديد لكن كوليسترولها عالٍ؛ إذا كوليسترولك مرتفع استشر طبيبك قبل الإكثار منها.

💡 نصائح:
• تجنب الشاي والقهوة أثناء الأكل مباشرة لأنها تقلل امتصاص الحديد.
• استخدم فلتر «الغنية بالحديد» في شاشة إضافة الوجبات، وتابع عمود الكوليسترول لكل وجبة.
""")
        txt.config(state="disabled")

    def open_plan(self):
        win, card = self.new_window("خطة الوجبات الأسبوعية 📅", 490, 580)
        lbl(card, "خطة رفع الحديد وخفض الكوليسترول", 11, True, C["GOLD"]).pack(pady=(0, 8))
        row = tk.Frame(card, bg=C["CARD_BG"])
        row.pack(fill="x", pady=(0, 8))
        txt = tk.Text(card, font=(FONT, 10), bg=C["MAIN_BG"], fg=C["TEXT"], bd=0, padx=10, pady=10, wrap="word")
        buttons = {}

        def show(i):
            for j, b in buttons.items():
                b.config(bg=C["ACCENT"] if j == i else C["MAIN_BG"], fg="white" if j == i else C["TEXT"])
            f, l, d, s = WEEK_PLAN[i]
            txt.config(state="normal")
            txt.delete("1.0", tk.END)
            txt.insert(tk.END, f"📅 {DAY_NAMES[i]}\n\n🥣 الفطور:\n{f}\n\n🥗 الغداء:\n{l}\n\n"
                               f"🍗 العشاء:\n{d}\n\n🍎 السناك:\n{s}\n\n"
                               "💡 الخطة إرشادية وليست بديلاً عن استشارة أخصائي التغذية.")
            txt.config(state="disabled")

        for i in DAY_ORDER:
            b = btn(row, DAY_NAMES[i], lambda i=i: show(i), C["MAIN_BG"], fg=C["TEXT"], size=8)
            b.pack(side="right", expand=True, fill="x", padx=1)
            buttons[i] = b
        txt.pack(fill="both", expand=True)
        show(datetime.now().weekday())

    # ---------- وصفة مركبة ----------
    def open_recipe(self, on_done=None):
        win, card = self.new_window("وصفة مركبة 🍳", 540, 800, True)
        lbl(card, "حوّل وصفتك لصنف جاهز (يتحسب لكل 100 جرام)", 11, True, C["GOLD"]).pack(pady=(0, 6))
        lbl(card, "اسم الوصفة:", 8, fg=C["MUTED"]).pack(anchor="e")
        e_name = ent(card)
        e_name.pack(fill="x", pady=(1, 6), ipady=3)
        lbl(card, "🔍 ابحث عن مكوّن:", 8, fg=C["MUTED"]).pack(anchor="e")
        e_search = ent(card)
        e_search.pack(fill="x", pady=(1, 4), ipady=3)
        lb = tk.Listbox(card, font=(FONT, 9), bg=C["MAIN_BG"], fg=C["TEXT"], selectbackground=C["ACCENT"],
                        selectforeground="white", bd=0, highlightthickness=0, height=5, justify="right",
                        exportselection=False)
        lb.pack(fill="x", pady=(0, 4))
        gf = tk.Frame(card, bg=C["CARD_BG"])
        gf.pack(fill="x", pady=(0, 4))
        lbl(gf, "الوزن الخام (جرام):", 8, fg=C["MUTED"]).pack(side="right")
        e_g = ent(gf, justify="center", width=8)
        e_g.insert(0, "100")
        e_g.pack(side="right", padx=6, ipady=3)
        ings = []

        def fill_list(*_):
            lb.delete(0, tk.END)
            q = norm(e_search.get())
            for n in sorted(self.foods()):
                if q in norm(n):
                    lb.insert(tk.END, n)

        def sums():
            t = [0.0] * 7
            for food, g in ings:
                v = self.get_food(food)
                for i in (0, 1, 2, 3, 4, 6):
                    t[i] += v[i] * g / 100
            return t, sum(g for _, g in ings)

        def render():
            for w in inner.winfo_children():
                w.destroy()
            for i, (food, g) in enumerate(ings):
                r = tk.Frame(inner, bg=C["MAIN_BG"], padx=8, pady=4)
                r.pack(fill="x", pady=2)
                lbl(r, f"{food} — {fmt(g)}ج", 9, bg=C["MAIN_BG"]).pack(side="right")
                btn(r, "🗑", lambda i=i: (ings.pop(i), render()), C["DANGER"]).pack(side="left")
            t, raw = sums()
            lbl_tot.config(text=f"المجموع الخام: {round(t[0])} kcal | {fmt(raw)} جرام | حديد {fmt(t[4])}mg | كوليسترول {round(t[6])}mg")

        def add_ing():
            sel = lb.curselection()
            if not sel:
                messagebox.showwarning("تنبيه", "اختر مكوّناً من القائمة.", parent=win)
                return
            try:
                g = float(e_g.get().replace(",", "."))
                if g <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("خطأ", "أدخل وزناً صحيحاً.", parent=win)
                return
            ings.append((lb.get(sel[0]), g))
            render()

        btn(gf, "➕ أضف المكوّن", add_ing, C["GOOD"], size=9).pack(side="left", ipadx=6, ipady=2)
        e_search.bind("<KeyRelease>", fill_list)
        hold = tk.Frame(card, bg=C["CARD_BG"], height=150)
        hold.pack(fill="x", pady=6)
        hold.pack_propagate(False)
        inner = make_scrollable(hold)
        lbl_tot = lbl(card, "", 8, True, C["GOLD"], wraplength=490)
        lbl_tot.pack(pady=(0, 6))
        lbl(card, "الوزن النهائي بعد الطبخ (جرام) — فاضي = مجموع الأوزان الخام:", 8, fg=C["MUTED"]).pack(anchor="e")
        e_final = ent(card, justify="center")
        e_final.pack(fill="x", pady=(1, 6), ipady=3)
        lbl(card, "عدد الحصص:", 8, fg=C["MUTED"]).pack(anchor="e")
        e_serv = ent(card, justify="center")
        e_serv.insert(0, "1")
        e_serv.pack(fill="x", pady=(1, 8), ipady=3)

        def save_recipe():
            name = e_name.get().strip()
            if not name or not ings:
                messagebox.showerror("خطأ", "اكتب اسم الوصفة وأضف مكوّناً واحداً على الأقل.", parent=win)
                return
            t, raw = sums()
            try:
                final = float(e_final.get().replace(",", ".")) if e_final.get().strip() else raw
                serv = float(e_serv.get().replace(",", ".")) if e_serv.get().strip() else 1
                if final <= 0 or serv <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("خطأ", "الوزن النهائي وعدد الحصص لازم يكونوا أرقاماً موجبة.", parent=win)
                return
            k = 100 / final
            self.d["custom_foods"][name] = [round(t[0] * k, 1), round(t[1] * k, 1), round(t[2] * k, 1),
                                            round(t[3] * k, 1), round(t[4] * k, 1),
                                            round(final / serv, 1), round(t[6] * k, 1)]
            self.save()
            win.destroy()
            if on_done:
                on_done()
            messagebox.showinfo("تم", f"تم حفظ ({name}) ضمن أكلاتك 🍲")

        btn(card, "💾 حفظ الوصفة", save_recipe, C["ACCENT"], size=10).pack(fill="x", ipady=7)
        fill_list()
        render()

    # ---------- المكملات ----------
    def open_supplements(self):
        win, card = self.new_window("المكملات 💊", 500, 580)
        lbl(card, "مكملاتي اليومية", 13, True, C["GOLD"]).pack(pady=(0, 8))
        af = tk.Frame(card, bg=C["CARD_BG"])
        af.pack(fill="x", pady=(0, 8))
        e_name = ent(af)
        e_name.pack(side="right", fill="x", expand=True, ipady=3)
        lbl(af, "الاسم / الجرعة:", 8, fg=C["MUTED"]).pack(side="right", padx=4)
        e_time = ent(af, justify="center", width=7)
        e_time.pack(side="left", ipady=3, padx=(0, 4))
        lbl(af, "تذكير (HH:MM):", 8, fg=C["MUTED"]).pack(side="left", padx=2)
        hold = tk.Frame(card, bg=C["CARD_BG"])
        keep = []

        def toggle(name, var):
            taken = self.today.setdefault("supp_taken", [])
            if var.get() and name not in taken:
                taken.append(name)
            elif not var.get() and name in taken:
                taken.remove(name)
            self.save()
            self.update_ui()

        def delete(name):
            self.d["supplements"] = [x for x in self.d["supplements"] if x["name"] != name]
            if name in self.today.get("supp_taken", []):
                self.today["supp_taken"].remove(name)
            self.save()
            self.update_ui()
            render()

        def add():
            name = e_name.get().strip()
            if not name:
                return
            tm = e_time.get().strip()
            if tm:
                mt = re.fullmatch(r"(\d{1,2}):(\d{2})", tm)
                if not mt or int(mt.group(1)) > 23 or int(mt.group(2)) > 59:
                    messagebox.showerror("خطأ", "الوقت لازم يكون بصيغة HH:MM مثل 09:00", parent=win)
                    return
                tm = f"{int(mt.group(1)):02d}:{mt.group(2)}"
            if any(x["name"] == name for x in self.d["supplements"]):
                messagebox.showerror("خطأ", "هالمكمل موجود من قبل.", parent=win)
                return
            self.d["supplements"].append({"name": name, "time": tm})
            self.save()
            e_name.delete(0, tk.END)
            e_time.delete(0, tk.END)
            self.update_ui()
            render()

        btn(card, "➕ أضف مكمل", add, C["GOOD"], size=9).pack(fill="x", ipady=4, pady=(0, 8))
        hold.pack(fill="both", expand=True)
        inner = make_scrollable(hold)

        def render():
            for w in inner.winfo_children():
                w.destroy()
            keep.clear()
            if not self.d["supplements"]:
                lbl(inner, "ما أضفت مكملات بعد.\nمثال: حديد 65mg ، فيتامين D", 10, fg=C["MUTED"],
                    justify="center").pack(pady=40)
                return
            for x in self.d["supplements"]:
                r = tk.Frame(inner, bg=C["MAIN_BG"], padx=10, pady=8)
                r.pack(fill="x", pady=3)
                var = tk.BooleanVar(value=x["name"] in self.today.get("supp_taken", []))
                keep.append(var)
                tk.Checkbutton(r, text=x["name"], variable=var, bg=C["MAIN_BG"], fg=C["TEXT"],
                               selectcolor=C["CARD_BG"], activebackground=C["MAIN_BG"], activeforeground=C["TEXT"],
                               font=(FONT, 10, "bold"),
                               command=lambda n=x["name"], v=var: toggle(n, v)).pack(side="right")
                if x.get("time"):
                    lbl(r, f"⏰ {x['time']}", 8, fg=C["MUTED"], bg=C["MAIN_BG"]).pack(side="right", padx=10)
                btn(r, "🗑", lambda n=x["name"]: delete(n), C["DANGER"]).pack(side="left")
        render()

    # ---------- الملخص الأسبوعي ----------
    def open_summary(self):
        win, card = self.new_window("الملخص الأسبوعي 📈", 560, 660)
        lbl(card, "أسبوعك مقارنة بالأسبوع اللي قبله", 12, True, C["GOLD"]).pack(pady=(0, 8))
        cur, prev = self.period_stats(0), self.period_stats(7)
        table = tk.Frame(card, bg=C["MAIN_BG"], padx=10, pady=10)
        table.pack(fill="x")
        for c, h in enumerate(("المؤشر", "هذا الأسبوع", "الأسبوع السابق", "التغيير")):
            lbl(table, h, 9, True, C["GOLD"], bg=C["MAIN_BG"]).grid(row=0, column=3 - c, padx=10, pady=(0, 6), sticky="e")
        rows = [("أيام التسجيل", "days", "", "up"), ("متوسط السعرات", "cal", " kcal", "neutral"),
                ("متوسط البروتين", "p", " ج", "up"), ("متوسط الحديد", "fe", " mg", "up"),
                ("متوسط الكوليسترول", "chol", " mg", "down"), ("متوسط الماء", "water", " لتر", "up"),
                ("أيام ضمن الهدف", "hit", "", "up")]
        for r, (name, key, unit, better) in enumerate(rows, start=1):
            a, b = cur[key], prev[key]
            diff = a - b
            if not prev["days"] or abs(diff) < 1e-9:
                chg, col = "—", C["MUTED"]
            else:
                good = (diff > 0 and better == "up") or (diff < 0 and better == "down")
                bad = (diff < 0 and better == "up") or (diff > 0 and better == "down")
                col = C["GOOD"] if good else C["DANGER"] if bad else C["MUTED"]
                chg = f"{'▲' if diff > 0 else '▼'} {fmt(abs(diff))}"
            f1 = (lambda v: f"{v:.1f}") if key == "water" else fmt
            for c, val in enumerate((name, f"{f1(a)}{unit}", f"{f1(b)}{unit}", chg)):
                lbl(table, val, 9, c == 0, col if c == 3 else C["TEXT"], bg=C["MAIN_BG"]).grid(
                    row=r, column=3 - c, padx=10, pady=3, sticky="e")

        cutoff = (datetime.now() - timedelta(days=14)).strftime("%Y-%m-%d")
        ws = [x for x in self.d["weights"] if x["date"] >= cutoff]
        if len(ws) >= 2:
            lbl(card, f"⚖️ الوزن: {ws[0]['weight']} ← {ws[-1]['weight']} كجم ({ws[-1]['weight'] - ws[0]['weight']:+.1f})",
                10, True, C["BLUE"]).pack(pady=(10, 0))

        tips = []
        if cur["days"] < 5:
            tips.append(f"📝 سجّلت {cur['days']} أيام فقط هالأسبوع؛ كل ما سجّلت أكتر كانت النتائج أدق.")
        if cur["days"] and cur["chol"] > self.chol_limit():
            tips.append(f"🫀 متوسط الكوليسترول ({round(cur['chol'])} mg) فوق حدك ({self.chol_limit()} mg). قلّل البيض والأحشاء والجمبري.")
        if cur["days"] and cur["fe"] < self.iron_need():
            tips.append(f"🩸 متوسط الحديد ({fmt(cur['fe'])} mg) أقل من احتياجك ({self.iron_need()} mg). زِد العدس والسبانخ مع الليمون.")
        if cur["water"] and cur["water"] < self.d["water_goal"] * 0.8:
            tips.append("💧 الماء أقل من هدفك؛ جرّب التذكير من الإعدادات.")
        if cur["days"] and cur["hit"] >= max(1, round(cur["days"] * 0.6)):
            tips.append("🎯 ممتاز! أغلب أيامك كانت ضمن هدف السعرات.")
        if not tips:
            tips.append("✅ أسبوع متوازن، كمّل!")
        tf = tk.Frame(card, bg=C["MAIN_BG"], padx=12, pady=10)
        tf.pack(fill="x", pady=12)
        lbl(tf, "ملاحظات", 10, True, C["GOLD"], bg=C["MAIN_BG"]).pack(anchor="e")
        for t in tips:
            lbl(tf, t, 9, bg=C["MAIN_BG"], wraplength=480, justify="right").pack(anchor="e", pady=3)

    # ---------- الإعدادات ----------
    def open_settings(self):
        win, card = self.new_window("الإعدادات ⚙️", 520, 780, True)
        s = self.d["settings"]
        lbl(card, "الإعدادات", 14, True, C["GOLD"]).pack(pady=(0, 4))

        def section(title):
            lbl(card, title, 10, True, C["GOLD"]).pack(anchor="e", pady=(10, 3))

        def row(text, widget_maker):
            f = tk.Frame(card, bg=C["CARD_BG"])
            f.pack(fill="x", pady=2)
            lbl(f, text, 9, fg=C["MUTED"]).pack(side="right")
            w = widget_maker(f)
            w.pack(side="right", padx=8, ipady=2)
            return w

        section("💧 الماء")
        e_wg = row("هدف الماء اليومي (لتر):", lambda f: ent(f, justify="center", width=7))
        e_wg.insert(0, str(self.d["water_goal"]))
        choices = {"كل ساعة": 1.0, "كل ساعتين": 2.0, "كل 3 ساعات": 3.0, "إيقاف": 0.0}
        cur_label = next((k for k, v in choices.items() if v == float(s.get("water_reminder_hrs", 2))), "كل ساعتين")
        rem_var = tk.StringVar(value=cur_label)
        row("تذكير الماء:", lambda f: optmenu(f, rem_var, *choices.keys()))

        section("🥩 توزيع الماكروز (٪ من السعرات — المجموع 100)")
        e_p = row("بروتين ٪:", lambda f: ent(f, justify="center", width=7))
        e_c = row("كارب ٪:", lambda f: ent(f, justify="center", width=7))
        e_f = row("دهون ٪:", lambda f: ent(f, justify="center", width=7))
        e_p.insert(0, str(s.get("p_pct", 25)))
        e_c.insert(0, str(s.get("c_pct", 50)))
        e_f.insert(0, str(s.get("f_pct", 25)))

        section("💡 التلميحات")
        vit_var = tk.BooleanVar(value=bool(s.get("vitc_tip", True)))
        tk.Checkbutton(card, text="نبّهني لإضافة فيتامين C مع وجبات الحديد", variable=vit_var,
                       bg=C["CARD_BG"], fg=C["TEXT"], selectcolor=C["MAIN_BG"], activebackground=C["CARD_BG"],
                       activeforeground=C["TEXT"], font=(FONT, 9)).pack(anchor="e")

        def save_settings():
            try:
                wg = float(e_wg.get().replace(",", "."))
                p, c, f = int(e_p.get()), int(e_c.get()), int(e_f.get())
                if wg <= 0 or min(p, c, f) < 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("خطأ", "تأكد من القيم (أرقام موجبة).", parent=win)
                return
            if p + c + f != 100:
                messagebox.showerror("خطأ", f"مجموع النسب لازم يكون 100 (الحالي {p + c + f}).", parent=win)
                return
            self.d["water_goal"] = wg
            s.update({"water_reminder_hrs": choices[rem_var.get()], "p_pct": p, "c_pct": c, "f_pct": f,
                      "vitc_tip": bool(vit_var.get())})
            self.save()
            self.update_ui()
            win.destroy()

        btn(card, "💾 حفظ الإعدادات", save_settings, C["ACCENT"], size=10).pack(fill="x", ipady=6, pady=(10, 4))

        section("💾 النسخ الاحتياطي (نسخة تلقائية يومياً بمجلد backups، آخر 7)")
        bf = tk.Frame(card, bg=C["CARD_BG"])
        bf.pack(fill="x")

        def now_backup():
            p = self.make_backup()
            if p:
                messagebox.showinfo("تم", f"تم حفظ النسخة:\n{p}", parent=win)
            else:
                messagebox.showerror("خطأ", "تعذّر حفظ النسخة.", parent=win)

        btn(bf, "💾 نسخ احتياطي الآن", now_backup, C["GOOD"]).pack(side="right", expand=True, fill="x", padx=2, ipady=5)
        btn(bf, "♻️ استعادة نسخة", lambda: self.restore_backup(win), C["BLUE"], fg=C["MAIN_BG"]).pack(
            side="right", expand=True, fill="x", padx=2, ipady=5)

        section("🧹 أخرى")

        def do_reset():
            win.destroy()
            self.reset_day()

        btn(card, "🔄 تصفير اليوم والبدء من جديد", do_reset, C["DANGER"]).pack(fill="x", ipady=5)

    # ---------- التصدير ----------
    def stamp(self):
        return datetime.now().strftime("%Y-%m-%d_%H-%M")

    def all_rows(self):
        days = dict(self.d["history"])
        days[self.today["date"]] = self.today
        rows = []
        for ds in sorted(days):
            for m in days[ds].get("meals", []):
                rows.append([ds, m.get("type", ""), m["name"], m["cal"], m.get("p", 0), m.get("c", 0),
                             m.get("f", 0), m.get("fe", 0), m.get("chol", 0), m.get("note", "")])
        return rows, days

    HEADER = ["التاريخ", "النوع", "الوجبة", "السعرات", "بروتين (ج)", "كارب (ج)", "دهون (ج)",
              "حديد (mg)", "كوليسترول (mg)", "ملاحظة"]

    def open_export(self):
        win, card = self.new_window("تصدير التقرير 📤", 420, 380)
        lbl(card, "اختر صيغة التصدير", 12, True, C["GOLD"]).pack(pady=(0, 10))
        for text, cmd, bg in (
            ("📄 ملف نصي TXT (تقرير اليوم)", self.export_txt, C["GOLD"]),
            ("📑 CSV (يفتح في Excel - كل الأيام)", self.export_csv, C["GOOD"]),
            ("📗 Excel XLSX (كل الأيام + ملخص)", self.export_xlsx, C["GOOD"]),
            ("🖨 تقرير HTML (اطبعه PDF من المتصفح)", self.export_html, C["ACCENT"]),
        ):
            btn(card, text, cmd, bg, fg="black" if bg == C["GOLD"] else "white", size=10).pack(fill="x", ipady=7, pady=4)
        lbl(card, "للـ PDF: افتح تقرير HTML ثم Ctrl+P ← «حفظ كـ PDF».\nيدعم العربي بشكل ممتاز.", 8,
            fg=C["MUTED"], justify="center").pack(pady=8)

    def _done(self, path):
        messagebox.showinfo("تم التصدير", f"تم حفظ الملف:\n{path}")
        open_path(path)

    def export_txt(self):
        if not self.meals:
            messagebox.showwarning("تنبيه", "لا توجد وجبات مسجلة اليوم!")
            return
        t = totals(self.meals)
        path = os.path.join(REPORT_DIR, f"CalorieMate_Report_{self.stamp()}.txt")
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write("=== تقرير اليوم - CalorieMate Pro ===\n")
                f.write(f"التاريخ: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n--- الوجبات ---\n")
                for i, m in enumerate(self.meals, 1):
                    f.write(f"{i}. [{m.get('type', '')}] {m['name']} - {m['cal']} سعرة | بروتين: {m.get('p', 0)}ج | "
                            f"كارب: {m.get('c', 0)}ج | دهون: {m.get('f', 0)}ج | حديد: {m.get('fe', 0)}mg | "
                            f"كوليسترول: {m.get('chol', 0)}mg"
                            + (f" | ملاحظة: {m['note']}" if m.get("note") else "") + "\n")
                f.write("\n--- الملخص ---\n")
                f.write(f"السعرات: {round(t['cal'])} / {int(self.d['daily_goal'])} kcal\n")
                f.write(f"البروتين: {round(t['p'], 1)} ج | الكارب: {round(t['c'], 1)} ج | الدهون: {round(t['f'], 1)} ج\n")
                f.write(f"الحديد: {round(t['fe'], 1)} / {self.iron_need()} ملغم\n")
                f.write(f"الكوليسترول: {round(t['chol'])} / {self.chol_limit()} ملغم\n")
                f.write(f"الماء: {self.today['water']:.2f} / {self.d['water_goal']} لتر\n")
            self._done(path)
        except Exception as ex:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء التصدير: {ex}")

    def export_csv(self):
        rows, _ = self.all_rows()
        if not rows:
            messagebox.showwarning("تنبيه", "لا توجد بيانات للتصدير!")
            return
        path = os.path.join(REPORT_DIR, f"CalorieMate_{self.stamp()}.csv")
        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(self.HEADER)
                w.writerows(rows)
            self._done(path)
        except Exception as ex:
            messagebox.showerror("خطأ", f"حدث خطأ: {ex}")

    def export_xlsx(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            messagebox.showinfo("مطلوب مكتبة", "لتصدير Excel ثبّت المكتبة بهالأمر:\n\npip install openpyxl\n\n"
                                               "أو استخدم CSV (يفتح في Excel مباشرة).")
            return
        rows, days = self.all_rows()
        if not rows:
            messagebox.showwarning("تنبيه", "لا توجد بيانات للتصدير!")
            return
        path = os.path.join(REPORT_DIR, f"CalorieMate_{self.stamp()}.xlsx")
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "الوجبات"
            ws.sheet_view.rightToLeft = True
            ws.append(self.HEADER)
            for r in rows:
                ws.append(r)
            ws2 = wb.create_sheet("ملخص الأيام")
            ws2.sheet_view.rightToLeft = True
            ws2.append(["التاريخ", "السعرات", "بروتين", "كارب", "دهون", "حديد", "كوليسترول", "ماء (لتر)"])
            for ds in sorted(days):
                if not days[ds].get("meals") and not days[ds].get("water"):
                    continue
                t = totals(days[ds].get("meals", []))
                ws2.append([ds, round(t["cal"]), round(t["p"], 1), round(t["c"], 1), round(t["f"], 1),
                            round(t["fe"], 1), round(t["chol"]), days[ds].get("water", 0)])
            for sheet in (ws, ws2):
                for col in sheet.columns:
                    sheet.column_dimensions[col[0].column_letter].width = 20
            wb.save(path)
            self._done(path)
        except Exception as ex:
            messagebox.showerror("خطأ", f"حدث خطأ: {ex}")

    def export_html(self):
        e = html.escape
        t = totals(self.meals)
        rows = "".join(
            f"<tr><td>{e(m.get('type', ''))}</td><td>{e(m['name'])}</td><td>{m['cal']}</td><td>{m.get('p', 0)}</td>"
            f"<td>{m.get('c', 0)}</td><td>{m.get('f', 0)}</td><td>{m.get('fe', 0)}</td><td>{m.get('chol', 0)}</td><td>{e(m.get('note', ''))}</td></tr>"
            for m in self.meals) or "<tr><td colspan='9'>لا توجد وجبات اليوم</td></tr>"
        week = ""
        for i in range(6, -1, -1):
            d = datetime.now() - timedelta(days=i)
            dd = self.day_data(d.strftime("%Y-%m-%d"))
            wt = totals(dd["meals"])
            week += (f"<tr><td>{d.strftime('%Y-%m-%d')} {DAY_NAMES[d.weekday()]}</td><td>{round(wt['cal'])}</td>"
                     f"<td>{round(wt['fe'], 1)}</td><td>{round(wt['chol'])}</td><td>{dd.get('water', 0):.2f}</td></tr>")
        css = ("body{font-family:Tahoma,Arial,sans-serif;margin:30px;color:#1b2430}"
               "h1{color:#1a535c}h2{margin-top:28px;color:#1a535c}"
               "table{border-collapse:collapse;width:100%;margin-top:8px}"
               "th,td{border:1px solid #ccd;padding:6px 10px;text-align:right}"
               "th{background:#1a535c;color:#fff}.box{background:#f1f5f9;padding:12px;border-radius:8px}")
        doc = (f"<!DOCTYPE html><html dir='rtl' lang='ar'><head><meta charset='utf-8'>"
               f"<title>تقرير CalorieMate</title><style>{css}</style></head><body>"
               f"<h1>تقرير CalorieMate Pro 🍎</h1><p>التاريخ: {datetime.now().strftime('%Y-%m-%d %H:%M')}</p>"
               f"<div class='box'>السعرات: <b>{round(t['cal'])}</b> / {int(self.d['daily_goal'])} kcal &nbsp;|&nbsp; "
               f"الحديد: <b>{round(t['fe'], 1)}</b> / {self.iron_need()} mg &nbsp;|&nbsp; "
               f"الكوليسترول: <b>{round(t['chol'])}</b> / {self.chol_limit()} mg &nbsp;|&nbsp; "
               f"الماء: <b>{self.today['water']:.2f}</b> / {self.d['water_goal']} لتر</div>"
               f"<h2>وجبات اليوم</h2><table><tr><th>النوع</th><th>الوجبة</th><th>سعرات</th><th>بروتين</th>"
               f"<th>كارب</th><th>دهون</th><th>حديد</th><th>كوليسترول</th><th>ملاحظة</th></tr>{rows}</table>"
               f"<h2>آخر 7 أيام</h2><table><tr><th>اليوم</th><th>سعرات</th><th>حديد</th><th>كوليسترول</th>"
               f"<th>ماء</th></tr>{week}</table></body></html>")
        path = os.path.join(REPORT_DIR, f"CalorieMate_Report_{self.stamp()}.html")
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write(doc)
            self._done(path)
        except Exception as ex:
            messagebox.showerror("خطأ", f"حدث خطأ: {ex}")


if __name__ == "__main__":
    App().run()
